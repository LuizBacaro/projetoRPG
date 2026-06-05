#!/usr/bin/env python3
"""Gera docs/governanca-agentes-fluxos.xlsx a partir das tabelas de governança."""

from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "docs" / "governanca-agentes-fluxos.xlsx"

HEADER_FILL = PatternFill(start_color="1F4E79", end_color="1F4E79", fill_type="solid")
HEADER_FONT = Font(bold=True, color="FFFFFF")
WRAP = Alignment(wrap_text=True, vertical="top")


def write_sheet(ws, headers: list[str], rows: list[list]) -> None:
    ws.append(headers)
    for cell in ws[1]:
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    for row in rows:
        ws.append(row)
    for row in ws.iter_rows(min_row=2, max_row=ws.max_row):
        for cell in row:
            cell.alignment = WRAP
    for col_idx, header in enumerate(headers, 1):
        max_len = len(str(header))
        for row in rows:
            if col_idx - 1 < len(row):
                max_len = max(max_len, min(len(str(row[col_idx - 1])), 80))
        ws.column_dimensions[get_column_letter(col_idx)].width = min(max_len + 2, 50)
    ws.freeze_panes = "A2"


def main() -> None:
    wb = Workbook()
    wb.remove(wb.active)

    ws = wb.create_sheet("Camadas")
    write_sheet(
        ws,
        ["Nível", "O quê", "Caminho", "Função", "Ligação automática"],
        [
            [1, "Governança geral", "AGENTS.md", "Idioma, camadas, deploy, agentes", "Cursor: regra solicitável; Copilot: copilot-instructions.md"],
            [2, "Arquitetura (detalhe)", "docs/arquitetura-camadas-solid.md, arquitetura-multi-jogo.md, ports-repositorios-servicos.md", "SOLID, Auth Hub, ports", "Manual / @ no chat"],
            [2, "Requisitos funcionais", ".cursor/requisitos/{dnd35,dnd5e,tormenta,gurps}/", "RF por jogo (o quê)", "Rule requisitos-implementacao.mdc"],
            [2, "Instructions por domínio", ".github/instructions/*.md", "Backend, frontend, migrations", "Copilot applyTo; Cursor @ficheiro"],
            [3, "Rules Cursor", ".cursor/rules/*.mdc", "Resumo + gate", "Por glob ao editar ficheiros"],
            [4, "Skills", ".cursor/skills/, .github/skills/", "Playbooks sob demanda", "Cursor: skills do projeto; 1-2 por tarefa"],
            [5, "Agentes e prompts", ".github/agents/, .github/prompts/", "Personas e atalhos", "Principalmente GitHub Copilot"],
            [6, "Hook", ".github/hooks/operational-safety.json", "Guarda terminal/SQL", "Se hooks ativos no runtime"],
            [7, "CI", ".github/workflows/ci.yml", "Lint, audit-dnd35-js, E2E Playwright", "Push/PR"],
            [7, "ADRs", "docs/adr/", "Decisões irreversíveis (monólito, Tormenta, dnd5e, apps adiado)", "Manual / @ no chat"],
            ["—", "Tracking manual", "MELHORIAS.md, HISTORICO_EVOLUCAO.md", "Débito e histórico", "Solto — não ligado a agentes"],
        ],
    )

    ws = wb.create_sheet("Agentes")
    write_sheet(
        ws,
        ["Agente", "Ficheiro", "Quando usar", "Delega para (subagentes)", "user-invocable"],
        [
            ["Fullstack Orchestrator", "fullstack-orchestrator.agent.md", "Feature/bug FE+BE+BD", "Backend, Frontend, DB/Migrations, PostgreSQL DBA, JS Auditor", "sim"],
            ["Fullstack API Contract Orchestrator", "fullstack-api-contract-orchestrator.agent.md", "Payload, validação, compatibilidade API", "Backend, Frontend, DB/Migrations", "sim"],
            ["Backend FastAPI Specialist", "backend-fastapi-specialist.agent.md", "API, services, auth, testes", "—", "sim"],
            ["Frontend Arena Specialist", "frontend-arena-specialist.agent.md", "HTML/CSS/JS, getApiUrl, ?v=", "—", "sim"],
            ["Database and Migrations Specialist", "database-and-migrations-specialist.agent.md", "Alembic, models, seeds", "—", "sim"],
            ["PostgreSQL Database Administrator", "postgresql-database-administrator.agent.md", "Postgres em execução", "—", "sim"],
            ["JavaScript Project Auditor", "javascript-project-auditor.agent.md", "Auditoria ampla JS", "—", "sim"],
            ["RPG Requirements Analyst", "rpg-requirements-analyst.agent.md", "Livro/PDF → RF", "Explore*, PDF Explore, Backend, Frontend, DB, orchestrators", "sim"],
            ["PDF Explore Specialist", "pdf-explore-specialist.agent.md", "Extrair PDF/OCR", "—", "sim"],
        ],
    )

    ws = wb.create_sheet("Prompts")
    write_sheet(
        ws,
        ["Nome do prompt", "Ficheiro", "Agente ligado", "Copilot (slash)", "Cursor"],
        [
            ["Nova Feature Fullstack", "nova-feature-fullstack.prompt.md", "Fullstack Orchestrator", "/Nova Feature Fullstack", "@.github/prompts/... + pedido"],
            ["Correção Bug Fullstack", "correcao-bug-fullstack.prompt.md", "Fullstack Orchestrator", "/Correcao Bug Fullstack", "idem"],
            ["Refatoração Segura por Domínio", "refatoracao-segura-por-dominio.prompt.md", "Fullstack Orchestrator", "/Refatoracao Segura por Dominio", "idem"],
            ["Auditoria Pré-Merge Fullstack", "auditoria-pre-merge-fullstack.prompt.md", "Fullstack Orchestrator", "/Auditoria Pre-Merge Fullstack", "idem"],
            ["Migration Segura", "migration-segura.prompt.md", "Fullstack API Contract Orchestrator", "/Migration Segura", "idem"],
            ["Deploy Readiness Fullstack", "deploy-readiness-fullstack.prompt.md", "Fullstack Orchestrator", "/Deploy Readiness Fullstack", "idem"],
            ["Investigação Performance Fullstack", "investigacao-performance-fullstack.prompt.md", "Fullstack Orchestrator", "/Investigacao Performance Fullstack", "idem"],
            ["Levantamento Requisitos RPG", "levantamento-requisitos-rpg.prompt.md", "RPG Requirements Analyst", "/Levantamento Requisitos RPG", "idem"],
        ],
    )

    ws = wb.create_sheet("Skills")
    write_sheet(
        ws,
        ["Skill", "Pasta", "Carregar quando", "Não usar para"],
        [
            ["arena-ttrpg-architecture", ".cursor/skills/", "CORS, Vercel, Render, getApiUrl, 404 API", "Regras de ficha"],
            ["arena-producao-dados-neon-render", ".cursor/skills/", "DATABASE_URL, Neon, migrations prod, /health/live", "—"],
            ["dnd-spellcasting-conventions", ".cursor/skills/", "Slots, grimório D&D 3.5", "dnd5e / Tormenta"],
            ["class-progression-conventions", ".cursor/skills/", "BBA, resistências, progressão classe", "—"],
            ["tormenta-20-arena-arquitetura-e-regras", ".cursor/skills/", "Ficha Tormenta, MB, /tormenta/", "Outros jogos"],
            ["gurps-4e-requisitos-ficha-arena", ".cursor/skills/", "Ficha GURPS", "Outros jogos"],
            ["rpg-requirements-analysis", ".github/skills/", "Livro/PDF → RF", "Implementação direta"],
            ["pdf", ".github/skills/", "Extrair/OCR PDF", "—"],
            ["javascript-audit", ".github/skills/", "Auditoria ampla JS", "Bug pontual"],
            ["frontend-design-teste", ".github/skills/", "UI nova com design", "Bug de API"],
            ["postgresql", ".github/skills/", "Migração Oracle→.NET", "Neon do Arena"],
        ],
    )

    ws = wb.create_sheet("Rules")
    write_sheet(
        ws,
        ["Ficheiro", "Glob", "alwaysApply", "Aponta para"],
        [
            ["arquitetura-camadas.mdc", "backend/**/*", "false", "AGENTS.md + docs/arquitetura-camadas-solid.md"],
            ["requisitos-implementacao.mdc", ".cursor/requisitos/**/*", "false", "AGENTS.md + gate RF + instructions + skills"],
            ["backend-qualidade-ci.mdc", "backend/**/*.py", "false", "docs/normas-qualidade-backend-ci.md, make ci-backend-lint"],
        ],
    )

    ws = wb.create_sheet("Onde_esta")
    write_sheet(
        ws,
        ["Pergunta", "Resposta", "Caminho"],
        [
            ["Fonte principal?", "AGENTS.md", "raiz/AGENTS.md"],
            ["Requisitos de ficha?", "Por jogo", ".cursor/requisitos/"],
            ["Regras backend?", "Instructions", ".github/instructions/backend.instructions.md"],
            ["Atalho Cursor backend?", "Rule", ".cursor/rules/arquitetura-camadas.mdc"],
            ["Orquestrador?", "Agente Copilot", ".github/agents/fullstack-orchestrator.agent.md"],
            ["Atalho feature?", "Prompt", ".github/prompts/nova-feature-fullstack.prompt.md"],
            ["Conjuração D&D 3.5?", "Skill", ".cursor/skills/dnd-spellcasting-conventions/"],
            ["Deploy/CORS?", "Skill", ".cursor/skills/arena-ttrpg-architecture/"],
            ["Auth refresh/OAuth?", "Doc", "docs/auth-sessao-oauth.md"],
            ["Livros → dados/RF?", "Doc", "docs/livros-para-dados.md"],
            ["Matriz RF × código?", "Doc", "docs/requisitos-cobertura-matrix.md"],
            ["Auditoria JS dnd35?", "Doc + script", "docs/dnd35-js-auditoria-2026-06.md, scripts/audit-dnd35-js.sh"],
            ["E2E Playwright?", "Doc", "docs/e2e-playwright-arena.md"],
            ["Dados Neon?", "Skill", ".cursor/skills/arena-producao-dados-neon-render/"],
            ["Guarda terminal?", "Hook", ".github/hooks/"],
            ["Guia markdown?", "Este doc", "docs/governanca-agentes-fluxos.md"],
            ["Planilha?", "Excel", "docs/governanca-agentes-fluxos.xlsx"],
        ],
    )

    ws = wb.create_sheet("Solto_ou_fraco")
    write_sheet(
        ws,
        ["Item", "Situação", "Ação recomendada"],
        [
            ["GUIA_ORQUESTRACAO_AGENTES.md", "Paralelo a governanca-agentes-fluxos.md", "Manter cross-link; editar AGENTS.md como fonte"],
            [".github/agents/ no Cursor", "Não são slash commands nativos", "Usar @ ou colar prompt"],
            ["Skill postgresql (.github)", "É para .NET, não Arena", "Não usar para Neon"],
            ["MELHORIAS.md", "Manual", "Citar @ quando relevante"],
            ["Hook cwd absoluto", "Pode falhar noutro PC", "Atualizar operational-safety.json"],
            ["dnd5e RF vs código", "Pode divergir de dnd35", "Verificar README dnd5e antes de implementar"],
            ["~/.cursor/skills-cursor/", "Fora do repo", "Não versionado"],
        ],
    )

    ws = wb.create_sheet("Receitas")
    write_sheet(
        ws,
        ["Tipo de tarefa", "Passo 1", "Passo 2", "Passo 3", "Passo 4"],
        [
            ["Nova mecânica D&D 3.5", "AGENTS.md", "rpg-requirements-analysis (+ pdf)", "Escrever RF em .cursor/requisitos/dnd35/", "Implementar + skill domínio"],
            ["Bug ficha D&D 3.5", "RF específico", "skill magia/progressão se aplicável", "games/dnd35/", "make ci-backend-lint"],
            ["Erro só produção", "arena-ttrpg-architecture", "arena-producao-dados se BD", "—", "PRE_DEPLOY_CHECKLIST"],
            ["Migration", "Prompt Migration Segura", "API Contract Orchestrator", "PRE_DEPLOY + skill Neon", "CI"],
            ["Levantamento livro", "Levantamento Requisitos RPG / RPG Analyst", "pdf skill", "RF em requisitos/", "Handoff — não codar"],
            ["Auditoria JS dnd35", "docs/dnd35-js-auditoria-2026-06.md", "scripts/audit-dnd35-js.sh", "Fase 1 quick wins", "E2E + ESLint"],
            ["Auth/OAuth produção", "docs/auth-sessao-oauth.md", "Render env GOOGLE_OAUTH_*", "PRE_DEPLOY_CHECKLIST", "—"],
        ],
    )

    OUT.parent.mkdir(parents=True, exist_ok=True)
    wb.save(OUT)
    print(f"Gerado: {OUT}")


if __name__ == "__main__":
    main()
