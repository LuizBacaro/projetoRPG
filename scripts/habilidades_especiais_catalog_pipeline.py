"""
Pipeline do catálogo de Habilidades Especiais (D&D 3.5 / Arena TTRPG).

Lê a aba `Habilidades especiais` da planilha
`Características especiais_v2.xlsx` (ou equivalente) e gera um catálogo
canônico em JSON, acessível por `slug`. Esse catálogo é consumido pelo
backend para enriquecer a ficha com descrições oficiais de habilidades que
hoje aparecem apenas como texto solto.

Arquitetura (SRP):
- `HabilidadesWorkbookReader` lê a aba e devolve linhas cruas.
- `HabilidadesNormalizer` gera registros canônicos e emite `ParseWarning`s.
- `HabilidadesCatalogExporter` serializa JSON.
- `HabilidadesCatalogPipeline` orquestra.

Contrato do JSON:
    {
      "source": "<arquivo.xlsx>",
      "generated_at": "ISO-8601",
      "total_habilidades": <int>,
      "habilidades": [
        {
          "slug": "furia",
          "titulo": "Fúria",
          "descricao": "<texto>",
          "aliases": []
        },
        ...
      ]
    }
"""

from __future__ import annotations

from dataclasses import dataclass, field
import json
from pathlib import Path
import re
import sys
from typing import Any

from openpyxl import load_workbook

# Suporta uso como pacote (`from scripts.habilidades_... import ...`) e
# carregamento direto por caminho de arquivo (usado nos testes).
_SCRIPTS_DIR = Path(__file__).resolve().parent
if str(_SCRIPTS_DIR) not in sys.path:
    sys.path.insert(0, str(_SCRIPTS_DIR))

from catalog_common import (  # noqa: E402
    build_envelope,
    slug as _slug_common,
    texto_limpo as _texto_limpo_common,
)


SHEET_NAME = "Habilidades especiais"
HEADER_ROW = 3
REQUIRED_HEADERS = ("Título", "Descrição")


@dataclass(frozen=True)
class ParseWarning:
    titulo: str
    campo: str
    mensagem: str

    def to_dict(self) -> dict[str, Any]:
        return {"titulo": self.titulo, "campo": self.campo, "mensagem": self.mensagem}


@dataclass
class ParseResult:
    habilidades: list[dict[str, Any]] = field(default_factory=list)
    warnings: list[ParseWarning] = field(default_factory=list)
    source_path: str = ""


def slug_habilidade(titulo: str) -> str:
    """Gera slug estável para uma habilidade (delega ao `catalog_common`)."""
    return _slug_common(titulo)


# Backcompat: alias local delegando ao módulo comum.
_texto_limpo = _texto_limpo_common


def _normalizar_descricao(valor: Any) -> str:
    texto = _texto_limpo(valor)
    if not texto:
        return ""
    # Preserva quebras de linha intencionais mas normaliza espaços em excesso
    linhas = [re.sub(r"\s+", " ", parte).strip() for parte in texto.splitlines()]
    return "\n".join(linha for linha in linhas if linha)


class HabilidadesWorkbookReader:
    """Lê a aba `Habilidades especiais` e devolve pares título/descrição."""

    def __init__(self, planilha_path: Path) -> None:
        self.planilha_path = planilha_path

    def read(self) -> tuple[list[dict[str, Any]], list[ParseWarning]]:
        if not self.planilha_path.exists():
            raise FileNotFoundError(
                f"Planilha não encontrada: {self.planilha_path}"
            )
        wb = load_workbook(self.planilha_path, data_only=True)
        if SHEET_NAME not in wb.sheetnames:
            raise ValueError(
                f"Aba {SHEET_NAME!r} não encontrada em {self.planilha_path.name}. "
                f"Abas disponíveis: {wb.sheetnames}"
            )

        ws = wb[SHEET_NAME]
        headers = [_texto_limpo(c.value) for c in ws[HEADER_ROW]]
        idx = {nome: i for i, nome in enumerate(headers) if nome}

        for obrigatorio in REQUIRED_HEADERS:
            if obrigatorio not in idx:
                raise ValueError(
                    f"Cabeçalho obrigatório ausente em "
                    f"{self.planilha_path.name}: {obrigatorio!r} "
                    f"(linha {HEADER_ROW})."
                )

        rows: list[dict[str, Any]] = []
        warnings: list[ParseWarning] = []
        for row_number in range(HEADER_ROW + 1, ws.max_row + 1):
            row = [c.value for c in ws[row_number]]
            valores = {
                nome: (row[pos] if pos < len(row) else None)
                for nome, pos in idx.items()
            }
            titulo = _texto_limpo(valores.get("Título"))
            if not titulo:
                continue
            descricao = _normalizar_descricao(valores.get("Descrição"))
            if not descricao:
                warnings.append(
                    ParseWarning(
                        titulo=titulo,
                        campo="descricao",
                        mensagem="Habilidade sem descrição na planilha",
                    )
                )
            rows.append({"titulo": titulo, "descricao": descricao})
        return rows, warnings


class HabilidadesNormalizer:
    """Converte linhas em registros canônicos, deduplicando por slug."""

    def normalize(self, rows: list[dict[str, Any]]) -> ParseResult:
        resultado = ParseResult()
        vistos: dict[str, int] = {}

        for row in rows:
            titulo = _texto_limpo(row.get("titulo"))
            if not titulo:
                continue
            slug = slug_habilidade(titulo)
            if not slug:
                resultado.warnings.append(
                    ParseWarning(
                        titulo=titulo,
                        campo="slug",
                        mensagem="Não foi possível gerar slug válido",
                    )
                )
                continue
            if slug in vistos:
                resultado.warnings.append(
                    ParseWarning(
                        titulo=titulo,
                        campo="slug",
                        mensagem=f"Slug duplicado ({slug!r}); mantendo primeira ocorrência.",
                    )
                )
                continue
            vistos[slug] = len(resultado.habilidades)
            resultado.habilidades.append(
                {
                    "slug": slug,
                    "titulo": titulo,
                    "descricao": row.get("descricao") or "",
                    "aliases": [],
                }
            )
        return resultado


class HabilidadesCatalogExporter:
    def __init__(self, output_path: Path) -> None:
        self.output_path = output_path

    def export(self, result: ParseResult) -> Path:
        payload = build_envelope(
            source_path=result.source_path,
            items_key="habilidades",
            items=result.habilidades,
        )
        self.output_path.parent.mkdir(parents=True, exist_ok=True)
        self.output_path.write_text(
            json.dumps(payload, ensure_ascii=False, indent=2),
            encoding="utf-8",
        )
        return self.output_path


class HabilidadesCatalogPipeline:
    def __init__(
        self,
        reader: HabilidadesWorkbookReader,
        normalizer: HabilidadesNormalizer,
        exporter: HabilidadesCatalogExporter | None = None,
    ) -> None:
        self.reader = reader
        self.normalizer = normalizer
        self.exporter = exporter

    def run(self) -> ParseResult:
        rows, warnings_leitura = self.reader.read()
        resultado = self.normalizer.normalize(rows)
        resultado.warnings = [*warnings_leitura, *resultado.warnings]
        resultado.source_path = str(self.reader.planilha_path)
        if self.exporter is not None:
            self.exporter.export(resultado)
        return resultado


__all__ = [
    "HEADER_ROW",
    "HabilidadesCatalogExporter",
    "HabilidadesCatalogPipeline",
    "HabilidadesNormalizer",
    "HabilidadesWorkbookReader",
    "ParseResult",
    "ParseWarning",
    "REQUIRED_HEADERS",
    "SHEET_NAME",
    "slug_habilidade",
]
