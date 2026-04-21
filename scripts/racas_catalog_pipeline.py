"""
Pipeline de catálogo de raças (D&D 3.5 / Arena TTRPG).

Objetivo:
- Ler a aba `Raças` das planilhas `Características especiais.xlsx` e
  `Características especiais_v2.xlsx`.
- Normalizar campos textuais (idiomas, talentos, habilidades, etc.).
- Parsear modificadores de habilidade de forma tolerante a variações de
  delimitação (`;`, `,`, quebras de linha).
- Emitir avisos estruturados (`ParseWarning`) quando algum token não for
  reconhecido, sem falhar silenciosamente.
- Exportar JSON estável compatível com o contrato atual em
  `docs/dados/racas_caracteristicas_catalogo.json`.

Arquitetura (SRP):
- `RacasWorkbookReader` apenas lê o workbook e retorna linhas cruas.
- `RacasNormalizer` aplica parsing e produz `ParseResult` com warnings.
- `RacasCatalogExporter` serializa para disco.
- `RacasCatalogPipeline` orquestra.

Compatibilidade:
- O shape do JSON de saída permanece idêntico ao gerado pela versão anterior
  do script para não exigir alterações no runtime do backend.
"""

from __future__ import annotations

from dataclasses import dataclass, field
import json
from pathlib import Path
import re
import sys
from typing import Any

from openpyxl import load_workbook

# Suporta tanto uso como pacote (`from scripts.racas_catalog_pipeline import ...`)
# quanto carregamento direto por caminho de arquivo (usado nos testes).
_SCRIPTS_DIR = Path(__file__).resolve().parent
if str(_SCRIPTS_DIR) not in sys.path:
    sys.path.insert(0, str(_SCRIPTS_DIR))

from catalog_common import (  # noqa: E402
    build_envelope,
    sem_acentos as _sem_acentos_common,
    slug as _slug_common,
    split_tokens,
    texto_limpo as _texto_limpo_common,
)


SHEET_NAME = "Raças"
HEADER_ROW = 3
REQUIRED_HEADERS = ("Raça",)
OPTIONAL_HEADERS = (
    "Modificadores de habilidades",
    "Tamanho",
    "Deslocamento",
    "Idiomas iniciais",
    "Talentos especiais",
    "Habilidades especiais",
    "Resistências",
    "Modificadores de ataque",
    "Modificadores de defesa",
    "Modificadores de perícia",
    "Classe favorecida",
)

_ATTR_ALIASES = {
    "FORCA": "forca",
    "DESTREZA": "destreza",
    "CONSTITUICAO": "constituicao",
    "INTELIGENCIA": "inteligencia",
    "SABEDORIA": "sabedoria",
    "CARISMA": "carisma",
}


@dataclass(frozen=True)
class ParseWarning:
    """Aviso estruturado de qualidade de dados da planilha."""

    raca: str
    campo: str
    mensagem: str
    token: str | None = None

    def to_dict(self) -> dict[str, Any]:
        return {
            "raca": self.raca,
            "campo": self.campo,
            "mensagem": self.mensagem,
            "token": self.token,
        }


@dataclass
class ParseResult:
    racas: list[dict[str, Any]] = field(default_factory=list)
    warnings: list[ParseWarning] = field(default_factory=list)
    source_path: str = ""


# Backcompat: aliases locais delegando ao módulo comum.
_sem_acentos = _sem_acentos_common
_slug = _slug_common
_texto_limpo = _texto_limpo_common


def _lista_linhas(valor: Any) -> list[str]:
    """Divide texto multi-linha em lista, removendo bullets comuns (`-`, `+`)."""
    texto = _texto_limpo(valor)
    if not texto or texto == "-":
        return []
    linhas: list[str] = []
    for parte in texto.replace("\r", "\n").split("\n"):
        item = parte.strip()
        if not item:
            continue
        while item and item[0] in {"-", "+", "*", "•"}:
            item = item[1:].strip()
        if item:
            linhas.append(item)
    return linhas


def _lista_idiomas(valor: Any) -> list[str]:
    """Idiomas aceitam `;`, `,` ou quebra de linha (via `catalog_common`)."""
    texto = _texto_limpo(valor)
    if not texto or texto == "-":
        return []
    return split_tokens(texto)


def _parse_deslocamento_metros(valor: Any) -> int | None:
    texto = _texto_limpo(valor).lower()
    if not texto:
        return None
    m = re.search(r"(\d+)", texto)
    return int(m.group(1)) if m else None


def _parse_modificadores_habilidade(
    valor: Any,
    raca: str,
) -> tuple[list[dict[str, Any]], list[ParseWarning]]:
    texto = _texto_limpo(valor)
    warnings: list[ParseWarning] = []
    if not texto or texto == "-":
        return [], warnings

    normal = _sem_acentos(texto).upper()
    tokens = split_tokens(normal)
    saida: list[dict[str, Any]] = []
    for token in tokens:
        m = re.match(r"([+-]\d+)\s+([A-Z ]+)$", token)
        if not m:
            warnings.append(
                ParseWarning(
                    raca=raca,
                    campo="modificadores_habilidade",
                    mensagem="Token não reconhecido (esperado '+N ATRIBUTO')",
                    token=token,
                )
            )
            continue
        valor_num = int(m.group(1))
        attr_raw = re.sub(r"\s+", " ", m.group(2).strip())
        attr = _ATTR_ALIASES.get(attr_raw)
        if not attr:
            warnings.append(
                ParseWarning(
                    raca=raca,
                    campo="modificadores_habilidade",
                    mensagem="Atributo desconhecido",
                    token=token,
                )
            )
            continue
        saida.append({"atributo": attr, "valor": valor_num})
    return saida, warnings


class RacasWorkbookReader:
    """Lê a aba `Raças` e devolve linhas cruas indexadas por cabeçalho."""

    def __init__(self, planilha_path: Path) -> None:
        self.planilha_path = planilha_path

    def read(self) -> tuple[list[dict[str, Any]], list[ParseWarning]]:
        if not self.planilha_path.exists():
            raise FileNotFoundError(
                f"Planilha não encontrada: {self.planilha_path}"
            )

        wb = load_workbook(self.planilha_path, data_only=True)
        if SHEET_NAME not in wb.sheetnames:
            raise ValueError(
                f"Aba {SHEET_NAME!r} não encontrada na planilha "
                f"{self.planilha_path.name}. Abas disponíveis: {wb.sheetnames}"
            )

        ws = wb[SHEET_NAME]
        headers = [_texto_limpo(c.value) for c in ws[HEADER_ROW]]
        idx = {nome: i for i, nome in enumerate(headers) if nome}

        warnings: list[ParseWarning] = []
        for obrigatorio in REQUIRED_HEADERS:
            if obrigatorio not in idx:
                raise ValueError(
                    f"Cabeçalho obrigatório ausente em {self.planilha_path.name}: "
                    f"{obrigatorio!r} (linha {HEADER_ROW})."
                )
        for opcional in OPTIONAL_HEADERS:
            if opcional not in idx:
                warnings.append(
                    ParseWarning(
                        raca="",
                        campo=opcional,
                        mensagem="Coluna opcional ausente no cabeçalho da aba Raças",
                    )
                )

        rows: list[dict[str, Any]] = []
        for row_number in range(HEADER_ROW + 1, ws.max_row + 1):
            row = [c.value for c in ws[row_number]]
            valores: dict[str, Any] = {
                nome: (row[pos] if pos < len(row) else None)
                for nome, pos in idx.items()
            }
            if not _texto_limpo(valores.get("Raça")):
                continue
            rows.append(valores)

        return rows, warnings


class RacasNormalizer:
    """Converte linhas cruas em registros canônicos, emitindo warnings."""

    def normalize(self, rows: list[dict[str, Any]]) -> ParseResult:
        resultado = ParseResult()

        for row in rows:
            nome = _texto_limpo(row.get("Raça"))
            if not nome:
                continue

            slug = _slug(nome)
            if not slug:
                resultado.warnings.append(
                    ParseWarning(
                        raca=nome,
                        campo="slug",
                        mensagem="Não foi possível gerar slug válido para a raça",
                    )
                )

            modificadores, warn_mods = _parse_modificadores_habilidade(
                row.get("Modificadores de habilidades"), nome
            )
            resultado.warnings.extend(warn_mods)

            registro = {
                "slug": slug,
                "nome": nome,
                "modificadores_habilidade": modificadores,
                "tamanho": _texto_limpo(row.get("Tamanho")),
                "deslocamento_metros": _parse_deslocamento_metros(
                    row.get("Deslocamento")
                ),
                "idiomas_iniciais": _lista_idiomas(row.get("Idiomas iniciais")),
                "talentos_especiais": _lista_linhas(row.get("Talentos especiais")),
                "habilidades_especiais": _lista_linhas(
                    row.get("Habilidades especiais")
                ),
                "resistencias": _lista_linhas(row.get("Resistências")),
                "modificadores_ataque": _lista_linhas(
                    row.get("Modificadores de ataque")
                ),
                "modificadores_defesa": _lista_linhas(
                    row.get("Modificadores de defesa")
                ),
                "modificadores_pericia": _lista_linhas(
                    row.get("Modificadores de perícia")
                ),
                "classe_favorecida": _texto_limpo(row.get("Classe favorecida")),
            }
            resultado.racas.append(registro)

        return resultado


class RacasCatalogExporter:
    """Serializa o catálogo canônico em JSON."""

    def __init__(self, output_path: Path) -> None:
        self.output_path = output_path

    def export(self, result: ParseResult) -> Path:
        payload = build_envelope(
            source_path=result.source_path,
            items_key="racas",
            items=result.racas,
        )
        self.output_path.parent.mkdir(parents=True, exist_ok=True)
        self.output_path.write_text(
            json.dumps(payload, ensure_ascii=False, indent=2),
            encoding="utf-8",
        )
        return self.output_path


class RacasCatalogPipeline:
    """Orquestra leitura, normalização e exportação do catálogo de raças."""

    def __init__(
        self,
        reader: RacasWorkbookReader,
        normalizer: RacasNormalizer,
        exporter: RacasCatalogExporter | None = None,
    ) -> None:
        self.reader = reader
        self.normalizer = normalizer
        self.exporter = exporter

    def run(self) -> ParseResult:
        rows, warnings_leitura = self.reader.read()
        resultado = self.normalizer.normalize(rows)
        resultado.warnings = [*warnings_leitura, *resultado.warnings]
        resultado.source_path = str(self.reader.planilha_path)
        if self.exporter is not None:
            self.exporter.export(resultado)
        return resultado


__all__ = [
    "ParseResult",
    "ParseWarning",
    "RacasCatalogExporter",
    "RacasCatalogPipeline",
    "RacasNormalizer",
    "RacasWorkbookReader",
    "SHEET_NAME",
    "HEADER_ROW",
    "REQUIRED_HEADERS",
    "OPTIONAL_HEADERS",
]
