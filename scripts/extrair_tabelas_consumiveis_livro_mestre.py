from __future__ import annotations

import subprocess
from pathlib import Path
import re

from openpyxl import Workbook


ROOT = Path(__file__).resolve().parents[1]
PDF = ROOT / "D&D 3.5 - Livro do Mestre.pdf"
TMP = ROOT / "docs" / "tmp_tabelas_consumiveis_livro_mestre.txt"
OUT = ROOT / "Tabela_7-17_7-23_7-24_Consumiveis.xlsx"


def extrair_texto() -> str:
    # Faixa ampliada para capturar a continuidade das tabelas 7-23 e 7-24.
    subprocess.run(
        [
            "pdftotext",
            "-layout",
            "-f",
            "230",
            "-l",
            "246",
            str(PDF),
            str(TMP),
        ],
        check=True,
    )
    return TMP.read_text(encoding="utf-8", errors="ignore")


def bloco(texto: str, inicio: str, fim: str) -> list[str]:
    i = texto.find(inicio)
    if i < 0:
        return []
    j = texto.find(fim, i + len(inicio))
    trecho = texto[i : (j if j >= 0 else len(texto))]
    linhas = [l.rstrip() for l in trecho.splitlines()]
    return [l for l in linhas if l.strip()]


def escrever_sheet(ws, linhas: list[str]) -> None:
    ws.append(["linha_extraida"])
    for linha in linhas:
        limpa = re.sub(r"[\x00-\x08\x0B-\x1F\x7F]", "", linha)
        ws.append([limpa])


def main() -> None:
    texto = extrair_texto()
    linhas_717 = bloco(texto, "Tabela 7–17: Poções e Óleos", "PERGAMINHOS")
    linhas_723 = bloco(texto, "Tabela 7–23: Pergaminhos de Magia", "Tabela 7–24: Pergaminhos de Magia")
    linhas_724 = bloco(texto, "Tabela 7–24: Pergaminhos de Magia", "VARINHAS")

    wb = Workbook()
    ws1 = wb.active
    ws1.title = "Tabela 7-17"
    escrever_sheet(ws1, linhas_717)

    ws2 = wb.create_sheet("Tabela 7-23")
    escrever_sheet(ws2, linhas_723)

    ws3 = wb.create_sheet("Tabela 7-24")
    escrever_sheet(ws3, linhas_724)

    wb.save(OUT)
    print(f"Arquivo gerado: {OUT}")


if __name__ == "__main__":
    main()
