"""
Reescreve a aba `Segmentado` das planilhas de classe com layout canônico:

  Nível | Bônus Base de Ataque | Fortitude | Reflexos | Vontade | Especial

Isso evita heurísticas que escolhiam linhas de dados (ex.: 3° com 6 colunas)
como "cabeçalho" e corrompiam o JSON (`tabelas_classes_catalogo.json`).

Fonte: texto em `ExtracaoBruta` coluna A (linhas com nível + progressão).

Cobre automaticamente **todas** as planilhas que seguem o padrão
`tabelas_classes_excel/Tabela_3-*_O_*.xlsx` (11 classes de progressão).

Uso (na raiz do repositório):
  python3 scripts/reestruturar_segmentado_tabelas_classes.py

Depois, regenere o JSON:
  python3 processar_tabelas_classes_excel.py
"""

from __future__ import annotations

import re
from pathlib import Path

from openpyxl import load_workbook

BASE_DIR = Path(__file__).resolve().parents[1]
INPUT_DIR = BASE_DIR / "tabelas_classes_excel"

def listar_planilhas_progressao_classe() -> list[Path]:
    """
    Todas as planilhas de progressão de classe (níveis 1–20, BBA/TR/Especial).

    Padrão de nome: `Tabela_3-*_O_*.xlsx` (ex.: O Bárbaro, O Bardo).

    Fora do escopo (outro layout / outro propósito): Magias Conhecidas (3-5, 3-10),
    Deuses (3-7), Dano desarmado monge (3-15), Inimigos prediletos (3-18).
    """
    if not INPUT_DIR.is_dir():
        return []
    return sorted(INPUT_DIR.glob("Tabela_3-*_O_*.xlsx"))

HEADERS = [
    "Nível",
    "Bônus Base de Ataque",
    "Fortitude",
    "Reflexos",
    "Vontade",
    "Especial",
]

_COL_LETTERS = ["C1", "C2", "C3", "C4", "C5", "C6", "C7"]


def _ler_linhas_extracao_bruta(path: Path) -> list[str]:
    wb = load_workbook(path, read_only=True, data_only=True)
    try:
        ws = wb["ExtracaoBruta"]
        linhas: list[str] = []
        for row in ws.iter_rows(min_row=2, values_only=True):
            if not row or row[0] is None:
                continue
            t = str(row[0]).strip()
            if t:
                linhas.append(t)
        return linhas
    finally:
        wb.close()


def _titulo_tabela(linhas: list[str]) -> str:
    for ln in linhas[:5]:
        if ln.strip().lower().startswith("tabela 3-"):
            return ln.strip()
    return ""


def _peel_quatro_tokens(rest: str) -> tuple[list[str], str] | None:
    """Após o marcador de nível: BBA, Fort, Ref, Vontade; resto = texto livre."""
    remainder = rest.strip()
    tokens: list[str] = []
    for _ in range(4):
        remainder = remainder.lstrip()
        if not remainder:
            return None
        m = re.match(r"^(\+\d+(?:/\+\d+)+)", remainder)
        if m:
            tokens.append(m.group(1))
            remainder = remainder[m.end() :]
            continue
        m = re.match(r"^(\+\d+)", remainder)
        if m:
            tokens.append(m.group(1))
            remainder = remainder[m.end() :]
            continue
        return None
    especial = re.sub(r"\s+", " ", remainder.strip())
    if not especial:
        especial = "-"
    return (tokens, especial)


def _parse_linha_nivel(linha: str) -> tuple[str, str, str, str, str, str] | None:
    """
    Aceita '12°', '12 °' ou linha do Monge '   1       +0 ...' (sem ° no nível 1–3 em alguns PDFs).
    """
    s = linha.strip()
    m = re.match(r"^(\d+)\s*°\s*(.*)$", s)
    if m:
        nivel = int(m.group(1))
        rest = m.group(2).strip()
    else:
        m2 = re.match(r"^(\d+)\s+(.*)$", s)
        if not m2:
            return None
        nivel = int(m2.group(1))
        rest = m2.group(2).strip()
        if not rest.startswith("+"):
            return None

    peeled = _peel_quatro_tokens(rest)
    if peeled is None:
        return None
    toks, especial = peeled
    nivel_txt = f"{nivel}°"
    return (nivel_txt, toks[0], toks[1], toks[2], toks[3], especial)


def _linha_eh_ruido(linha: str) -> bool:
    x = linha.strip().lower()
    if x.startswith("linha"):
        return True
    if x.startswith("tabela 3-"):
        return True
    if "nível de ataque fortitude" in x:
        return True
    # Cabeçalhos quebrados do PDF (Monge — linha com Rajada/Dano/Desarmado)
    if "rajada de golpes" in x and "desarmado" in x:
        return True
    if "bônus base" in x and "bônus de ataque da" in x and "dano" in x:
        return True
    return False


def _linha_eh_cabecalho_tabela_quebrado(linha: str) -> bool:
    """Linhas de PDF em que o cabeçalho da tabela foi parar no meio do texto."""
    x = linha.strip().lower()
    if "fortitude" in x and "reflexos" in x and "vontade" in x and "especial" in x:
        return True
    if "nível" in x and "ataque" in x and "fortitude" in x and len(x) > 50:
        return True
    if "bônus base" in x and "nível" in x and len(x) > 40:
        return True
    return False


def _reestruturar_workbook(path: Path) -> tuple[int, list[str]]:
    linhas = _ler_linhas_extracao_bruta(path)
    titulo = _titulo_tabela(linhas)
    buffer_flavor: list[str] = []
    saida: list[tuple[str, str, str, str, str, str]] = []

    for ln in linhas:
        if _linha_eh_ruido(ln):
            continue
        if "nível" in ln.lower() and "bônus" in ln.lower() and "fortitude" in ln.lower():
            continue
        if _linha_eh_cabecalho_tabela_quebrado(ln):
            buffer_flavor.clear()
            continue

        parsed = _parse_linha_nivel(ln)
        if parsed is None:
            low = ln.strip().lower()
            if (
                len(low) > 12
                and not re.match(r"^\d+", ln.strip())
                and not _linha_eh_cabecalho_tabela_quebrado(ln)
            ):
                buffer_flavor.append(ln.strip())
            continue

        nivel, bba, fort, ref, von, esp = parsed
        prefix = ""
        if buffer_flavor:
            prefix = " | ".join(buffer_flavor) + (" | " if esp and esp != "-" else "")
            buffer_flavor = []
        if esp == "-":
            merged = prefix.rstrip(" | ") if prefix else "-"
        else:
            merged = prefix + esp if prefix else esp
        if not merged.strip():
            merged = "-"
        saida.append((nivel, bba, fort, ref, von, merged))

    wb = load_workbook(path)
    try:
        if "Segmentado" not in wb.sheetnames:
            raise ValueError(f"Sem aba Segmentado: {path.name}")
        ws = wb["Segmentado"]
        ws.delete_rows(1, ws.max_row)

        for col, letter in enumerate(_COL_LETTERS, start=1):
            ws.cell(1, col, letter)
        ws.cell(2, 1, titulo or path.stem)
        for col, h in enumerate(HEADERS, start=1):
            ws.cell(3, col, h)
        r = 4
        for row in saida:
            for col, val in enumerate(row, start=1):
                ws.cell(r, col, val)
            r += 1
        wb.save(path)
    finally:
        wb.close()

    return len(saida), [f"{path.name}: {len(saida)} linhas de progressão"]


def verificar_segmentado_pos_escrita(path: Path, linhas_esperadas: int) -> list[str]:
    """Confere cabeçalho canônico e quantidade de níveis (20 para classes completas)."""
    avisos: list[str] = []
    wb = load_workbook(path, read_only=True, data_only=True)
    try:
        ws = wb["Segmentado"]
        if ws.cell(3, 1).value != HEADERS[0]:
            avisos.append(f"{path.name}: célula A3 deveria ser {HEADERS[0]!r}")
        if ws.cell(3, 2).value != HEADERS[1]:
            avisos.append(f"{path.name}: B3 deveria ser cabeçalho BBA")
        if linhas_esperadas < 19:
            avisos.append(f"{path.name}: só {linhas_esperadas} linhas de nível (esperado ~20)")
    finally:
        wb.close()
    return avisos


def main() -> int:
    if not INPUT_DIR.is_dir():
        raise SystemExit(f"Diretório não encontrado: {INPUT_DIR}")
    planilhas = listar_planilhas_progressao_classe()
    if not planilhas:
        raise SystemExit(f"Nenhuma planilha Tabela_3-*_O_*.xlsx em {INPUT_DIR}")

    total = 0
    todos_avisos: list[str] = []
    for p in planilhas:
        n, msgs = _reestruturar_workbook(p)
        total += n
        for m in msgs:
            print(m)
        todos_avisos.extend(verificar_segmentado_pos_escrita(p, n))

    if todos_avisos:
        for a in todos_avisos:
            print(f"⚠️  {a}")
    else:
        print("✅ Verificação pós-escrita: cabeçalhos e contagem OK em todas.")

    print(
        f"✅ Concluído: {len(planilhas)} planilhas de classe, {total} linhas de progressão no total."
    )
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
