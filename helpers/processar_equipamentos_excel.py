"""
Script para processar planilha Excel de equipamentos D&D 3.5 (Tabela 7-5).
Gera backend/scripts/seed_equipamentos.py com EQUIPAMENTOS_DADOS e seed_equipamentos().
Usa openpyxl (sem dependência de pandas).
"""

from __future__ import annotations

import re
from pathlib import Path

from openpyxl import load_workbook

PLANILHA_PATH = Path("Tabela_7-5_Armas_EQUIPAMENTOS.xlsx")
OUTPUT_PATH = Path("backend/scripts/seed_equipamentos.py")


def _cell(row: tuple, idx: int):
    if idx >= len(row):
        return None
    v = row[idx]
    if v is None:
        return None
    s = str(v).strip()
    return s if s else None


def processar_planilha_equipamentos() -> list[dict]:
    """Lê a planilha (cabeçalho na linha 4) e retorna lista de dicts compatível com o modelo Equipamento."""
    if not PLANILHA_PATH.exists():
        raise FileNotFoundError(f"Planilha não encontrada: {PLANILHA_PATH}")

    wb = load_workbook(PLANILHA_PATH, data_only=True)
    ws = wb.active

    header_row = 4
    rows = list(ws.iter_rows(min_row=header_row, max_row=header_row, values_only=True))[0]
    headers = [str(c).strip() if c is not None else "" for c in rows]

    # Índices esperados (mesma ordem do Excel atual)
    idx = {name: i for i, name in enumerate(headers)}

    def col(name: str) -> int:
        if name not in idx:
            raise KeyError(f"Coluna '{name}' não encontrada. Cabeçalhos: {headers}")
        return idx[name]

    equipamentos: list[dict] = []

    for row in ws.iter_rows(min_row=header_row + 1, values_only=True):
        arma = _cell(row, col("Arma"))
        if not arma:
            continue

        categoria = _cell(row, col("Categoria"))
        if categoria and re.search(r"^[¹²³⁴⁵]", categoria):
            continue
        if categoria and "Notas da tabela" in categoria:
            continue
        if re.search(r"^[¹²³⁴⁵]", arma):
            continue

        def norm(v: str | None, empty_as_none: bool = True) -> str | None:
            if v is None or v == "—":
                return None if empty_as_none else v
            return v

        # Colunas A–J da planilha (linha 4 = cabeçalho): sem coluna de página no arquivo atual.
        pagina_ref = None
        for nome_col in ("Página", "Pág.", "Pág", "Referência", "PHB"):
            if nome_col in idx:
                pagina_ref = norm(_cell(row, col(nome_col)))
                if pagina_ref:
                    break

        equipamento = {
            "nome": arma,
            "categoria": norm(categoria),
            "subcategoria": norm(_cell(row, col("Subcategoria"))),
            "custo": norm(_cell(row, col("Custo"))),
            "dano_pequeno": norm(_cell(row, col("Dano (P)"))),
            "dano_medio": norm(_cell(row, col("Dano (M)"))),
            "critico": norm(_cell(row, col("Crítico"))),
            "alcance_incremento": norm(_cell(row, col("Alcance / incremento"))),
            "peso": norm(_cell(row, col("Peso"))),
            "tipo_dano": norm(_cell(row, col("Tipo de dano"))),
            "pagina_referencia": pagina_ref,
            "ativo": True,
        }
        equipamentos.append(equipamento)

    print(f"📊 Processando {len(equipamentos)} equipamentos válidos")
    return equipamentos


def _py_repr_dict(d: dict) -> str:
    """Gera literal Python para um dict (strings com repr)."""
    lines = ["    {"]
    for key, value in d.items():
        if value is None:
            lines.append(f"        {key!r}: None,")
        elif isinstance(value, bool):
            lines.append(f"        {key!r}: {value},")
        elif isinstance(value, str):
            lines.append(f"        {key!r}: {value!r},")
        else:
            lines.append(f"        {key!r}: {value},")
    lines.append("    },")
    return "\n".join(lines)


def gerar_seed_script(equipamentos: list[dict]) -> str:
    blocos = "\n".join(_py_repr_dict(eq) for eq in equipamentos)

    # Template estático (sem f-string) para não confundir chaves com interpolação.
    return (
        '"""\n'
        "Seed de equipamentos D&D 3.5 — Tabela 7-5 (armas).\n"
        "Gerado automaticamente por processar_equipamentos_excel.py a partir da planilha Excel.\n"
        "Não editar EQUIPAMENTOS_DADOS à mão; regenere o arquivo com o script.\n"
        '"""\n\n'
        "from datetime import datetime, timezone\n\n"
        "EQUIPAMENTOS_DADOS = [\n"
        f"{blocos}\n"
        "]\n\n"
        r'''
def _montar_descricao(eq: dict) -> str | None:
    """Campo legado: a UI usa colunas estruturadas; evitar duplicar o que já está nos campos."""
    return None


def seed_equipamentos(db) -> None:
    """
    Sincroniza o catálogo com a Tabela 7-5.

    - Remove entradas legadas (sem categoria) não usadas em equipamentos_jogador.
    - Para cada item da planilha: insere ou atualiza por nome (entre ativos).
    """
    from app.games.dnd35.models.equipamento import Equipamento, EquipamentoJogador

    usados = {
        row[0]
        for row in db.query(EquipamentoJogador.equipamento_id).distinct().all()
        if row[0] is not None
    }

    legado = (
        db.query(Equipamento)
        .filter(Equipamento.deleted_at.is_(None), Equipamento.categoria.is_(None))
        .all()
    )
    removidos = 0
    for eq in legado:
        if eq.id not in usados:
            db.delete(eq)
            removidos += 1
    if removidos:
        db.flush()
        print(f"🧹 Removidos {removidos} equipamentos legados (sem categoria, não referenciados).")

    inseridos = 0
    atualizados = 0

    for data in EQUIPAMENTOS_DADOS:
        nome = data["nome"]
        descricao = _montar_descricao(data)
        row = (
            db.query(Equipamento)
            .filter(Equipamento.nome == nome, Equipamento.deleted_at.is_(None))
            .first()
        )
        campos = {
            "descricao": descricao,
            "categoria": data.get("categoria"),
            "subcategoria": data.get("subcategoria"),
            "custo": data.get("custo"),
            "dano_pequeno": data.get("dano_pequeno"),
            "dano_medio": data.get("dano_medio"),
            "critico": data.get("critico"),
            "alcance_incremento": data.get("alcance_incremento"),
            "peso": data.get("peso"),
            "tipo_dano": data.get("tipo_dano"),
            "pagina_referencia": data.get("pagina_referencia"),
            "ativo": data.get("ativo", True),
        }
        if row:
            for k, v in campos.items():
                setattr(row, k, v)
            atualizados += 1
        else:
            db.add(
                Equipamento(
                    nome=nome,
                    **campos,
                    criado_em=datetime.now(timezone.utc),
                )
            )
            inseridos += 1

    db.commit()
    try:
        from app.shared.core.config import settings
        from app.shared.core.catalog_cache import catalog_cache

        if settings.CACHE_ENABLED:
            catalog_cache.invalidate_prefix("equipamentos:")
    except Exception:
        pass
    print(
        f"✅ Catálogo Tabela 7-5 sincronizado: +{inseridos} inseridos, "
        f"{atualizados} atualizados (total planilha: {len(EQUIPAMENTOS_DADOS)})."
    )
'''
    )


if __name__ == "__main__":
    equipamentos = processar_planilha_equipamentos()
    if not equipamentos:
        raise SystemExit("Nenhum equipamento válido na planilha.")
    OUTPUT_PATH.parent.mkdir(parents=True, exist_ok=True)
    OUTPUT_PATH.write_text(gerar_seed_script(equipamentos), encoding="utf-8")
    print(f"✅ Gerado: {OUTPUT_PATH} ({len(equipamentos)} itens)")
