"""
Script para processar a planilha Excel da Tabela 7-6 (Armaduras e Escudos).
Gera um script de seed para o catálogo armaduras_protecao.
"""

from __future__ import annotations

from pathlib import Path
from typing import Any

from openpyxl import load_workbook


PLANILHA_PATH = Path("Tabela_7-6_Armaduras_e_Escudos_EQUIPAMENTOS.xlsx")
OUTPUT_PATH = Path("backend/scripts/seed_armaduras_protecao.py")


def _limpar_texto(valor: Any) -> str | None:
    if valor is None:
        return None
    texto = str(valor).strip()
    if not texto or texto in {"-", "—"}:
        return None
    return texto


def _parse_bonus_ca(texto: str | None) -> int:
    if not texto:
        return 0
    texto = texto.replace("+", "").strip()
    try:
        return int(texto)
    except ValueError:
        return 0


def _parse_penalidade(texto: str | None) -> tuple[int, str | None]:
    if not texto:
        return 0, None

    texto_normalizado = texto.strip()
    if texto_normalizado.lower() == "especial":
        return 0, "Penalidade especial (consultar descricao do item)."

    try:
        return int(texto_normalizado), None
    except ValueError:
        return 0, f"Penalidade nao numerica: {texto_normalizado}"


def _parse_peso_kg(texto: str | None) -> float | None:
    if not texto:
        return None

    sem_kg = texto.lower().replace("kg", "").replace("+", "").strip()
    sem_kg = sem_kg.replace(",", ".")
    try:
        return float(sem_kg)
    except ValueError:
        return None


def _mapear_tipo(categoria: str) -> str:
    categoria_lower = categoria.lower()
    if "escudo" in categoria_lower:
        return "Escudo"
    if "acessor" in categoria_lower:
        return "Acessorio"
    return "Armadura"


def carregar_armaduras_da_planilha() -> list[dict[str, Any]]:
    if not PLANILHA_PATH.exists():
        raise FileNotFoundError(f"Planilha nao encontrada: {PLANILHA_PATH}")

    wb = load_workbook(PLANILHA_PATH, data_only=True)
    ws = wb["Tabela 7-6"]

    # Cabecalho esperado:
    # Categoria | Item | Custo | Bonus Armadura/Escudo | Max Destreza |
    # Penalidade | Falha Magia Arcana | Deslocamento (9 m) | Deslocamento (6 m) | Peso
    itens: list[dict[str, Any]] = []
    for row in ws.iter_rows(min_row=2, values_only=True):
        categoria = _limpar_texto(row[0])
        nome = _limpar_texto(row[1])
        custo = _limpar_texto(row[2])
        bonus = _limpar_texto(row[3])
        des_max = _limpar_texto(row[4])
        penalidade_txt = _limpar_texto(row[5])
        falha_arcana = _limpar_texto(row[6])
        deslocamento_9 = _limpar_texto(row[7])
        deslocamento_6 = _limpar_texto(row[8])
        peso_txt = _limpar_texto(row[9])

        if not categoria or not nome:
            continue

        penalidade, observacao_penalidade = _parse_penalidade(penalidade_txt)
        deslocamento = None
        if deslocamento_9 or deslocamento_6:
            deslocamento = f"9m: {deslocamento_9 or '-'} | 6m: {deslocamento_6 or '-'}"

        observacoes = []
        if custo:
            observacoes.append(f"Custo: {custo}")
        if observacao_penalidade:
            observacoes.append(observacao_penalidade)
        if categoria.lower().startswith("acessor"):
            observacoes.append("Item acessorio da tabela de armaduras e escudos.")

        item = {
            "nome": nome,
            "tipo": _mapear_tipo(categoria),
            "bonus_ca": _parse_bonus_ca(bonus),
            "des_max": des_max,
            "penalidade": penalidade,
            "falha_arcana": falha_arcana,
            "deslocamento": deslocamento,
            "peso": _parse_peso_kg(peso_txt),
            "propriedades_especiais": " | ".join(observacoes) if observacoes else None,
            "ativo": True,
        }
        itens.append(item)

    return itens


def gerar_seed_script(itens: list[dict[str, Any]]) -> str:
    linhas_itens = []
    for item in itens:
        blocos = []
        for chave, valor in item.items():
            if isinstance(valor, str):
                blocos.append(f"        '{chave}': {valor!r},")
            else:
                blocos.append(f"        '{chave}': {valor},")
        item_str = "    {\n" + "\n".join(blocos) + "\n    },"
        linhas_itens.append(item_str)

    itens_formatados = "\n".join(linhas_itens)

    return f'''"""
Seed de armaduras e itens de protecao D&D 3.5 (Tabela 7-6).
Gerado automaticamente a partir da planilha Excel.
"""

from datetime import datetime, timezone

from app.games.dnd35.models.armadura_protecao import ArmaduraProtecao


ARMADURAS_PROTECAO_DADOS = [
{itens_formatados}
]


def seed_armaduras_protecao(db):
    """
    Popula o catalogo de armaduras/itens de protecao da Tabela 7-6.
    """
    count = db.query(ArmaduraProtecao).count()
    if count > 0:
        print(f"✅ Armaduras/Protecao ja existem ({{count}}). Pulando seed.")
        return

    print("📦 Iniciando seed de armaduras/protecao da Tabela 7-6...")
    for item_data in ARMADURAS_PROTECAO_DADOS:
        item = ArmaduraProtecao(
            nome=item_data["nome"],
            tipo=item_data["tipo"],
            bonus_ca=item_data["bonus_ca"],
            des_max=item_data.get("des_max"),
            penalidade=item_data["penalidade"],
            falha_arcana=item_data.get("falha_arcana"),
            deslocamento=item_data.get("deslocamento"),
            peso=item_data.get("peso"),
            propriedades_especiais=item_data.get("propriedades_especiais"),
            ativo=item_data.get("ativo", True),
            criado_em=datetime.now(timezone.utc),
        )
        db.add(item)

    db.commit()
    print(f"✅ {{len(ARMADURAS_PROTECAO_DADOS)}} itens inseridos com sucesso!")
'''


def main() -> None:
    itens = carregar_armaduras_da_planilha()
    if not itens:
        raise ValueError("Nenhum item valido encontrado na planilha.")

    conteudo = gerar_seed_script(itens)
    OUTPUT_PATH.parent.mkdir(parents=True, exist_ok=True)
    OUTPUT_PATH.write_text(conteudo, encoding="utf-8")

    print(f"✅ Script de seed gerado: {OUTPUT_PATH}")
    print(f"📊 {len(itens)} itens processados")
    print("🔍 Preview dos primeiros 5 itens:")
    for idx, item in enumerate(itens[:5], start=1):
        print(f"  {idx}. {item['nome']} ({item['tipo']}, CA +{item['bonus_ca']})")


if __name__ == "__main__":
    main()
