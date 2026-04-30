"""
Processa o arquivo Tabela_7-17_7-23_7-24_Consumiveis.xlsx e gera
backend/scripts/seed_consumiveis.py com dados normalizados de consumíveis.
"""

from __future__ import annotations

import re
from pathlib import Path

from openpyxl import load_workbook

PLANILHA_PATH = Path("Tabela_7-17_7-23_7-24_Consumiveis.xlsx")
OUTPUT_PATH = Path("backend/scripts/seed_consumiveis.py")

_RANGE = r"\d{1,3}\s*[–-]\s*\d{1,3}|\d{1,3}"
_PRECO = r"\d[\d\.\s]*PO(?:\s*e\s*\d+\s*PP)?"
_ENTRY_RE = re.compile(
    rf"(?P<faixa>{_RANGE})\s*(?P<nome>.*?)\s+(?P<preco>{_PRECO})",
    flags=re.IGNORECASE,
)


def _clean_spaces(text: str) -> str:
    text = (
        text.replace("\xa0", " ")
        .replace("", " ")
        .replace("", " ")
        .replace("—", "—")
    )
    return re.sub(r"\s+", " ", text).strip()


def _normalizar_nome(nome: str) -> str:
    nome = _clean_spaces(nome)
    nome = re.sub(r"^:?\s*Pergaminhos de Magia\s+", "", nome, flags=re.IGNORECASE)
    nome = re.sub(rf"^(?:[^\wà-ÿ]*)(?:{_RANGE}\s*)+", "", nome, flags=re.IGNORECASE)
    nome = re.sub(r"^[-–—]+\s*", "", nome)
    nome = re.sub(r"\s+\(poção\)\s+\(poção\)$", " (poção)", nome, flags=re.IGNORECASE)
    nome = re.sub(r"\s+\(óleo\)\s+\(óleo\)$", " (óleo)", nome, flags=re.IGNORECASE)
    nome = re.sub(r"^[-–—]+\s*", "", nome)
    nome = re.sub(r"\s+", " ", nome)
    while re.match(r"^[\d\s–-]+", nome):
        novo = re.sub(r"^[\d\s–-]+", "", nome).lstrip()
        if novo == nome:
            break
        nome = novo
    return nome.strip(" -–—")


def _extrair_faixa(raw_nome: str) -> str | None:
    raw = _clean_spaces(raw_nome)
    m = re.match(rf"^\s*(?P<faixa>{_RANGE})\b", raw, flags=re.IGNORECASE)
    if not m:
        return None
    faixa = _clean_spaces(m.group("faixa"))
    return faixa if faixa else None


def _nome_valido(nome: str) -> bool:
    n = nome.lower()
    if len(nome) < 3 or len(nome) > 90:
        return False
    if not re.search(r"[a-zà-ÿ]", n):
        return False
    if n.startswith("(") or n.startswith("/") or "— —" in n:
        return False
    invalid_tokens = [
        "capítulo",
        "itens mágicos",
        "forjar anel",
        "forjar bastão",
        "preço",
        "custo",
        "nc ",
        "anel ",
        "bastão ",
        "metamágico",
        "tabela 7",
        "magias de",
        "preço de mercado",
    ]
    if any(tok in n for tok in invalid_tokens):
        return False
    if re.match(r"^[\W_]+$", n):
        return False
    return True


def _inferir_categoria_tipo(sheet_title: str, nome: str) -> tuple[str, str]:
    lower = nome.lower()
    if sheet_title == "Tabela 7-17":
        if "(poção)" in lower:
            return "Consumível mágico", "Poção"
        if "(óleo)" in lower:
            return "Consumível mágico", "Óleo"
        return "Consumível mágico", "Poção/Óleo"
    if sheet_title == "Tabela 7-23":
        return "Pergaminho", "Arcana"
    return "Pergaminho", "Divina"


def extrair_consumiveis() -> list[dict]:
    if not PLANILHA_PATH.exists():
        raise FileNotFoundError(f"Planilha não encontrada: {PLANILHA_PATH}")

    wb = load_workbook(PLANILHA_PATH, data_only=True)
    itens: list[dict] = []
    vistos: set[tuple[str, str]] = set()

    for ws in wb.worksheets:
        for row in ws.iter_rows(min_row=2, values_only=True):
            raw = row[0]
            if not raw or not isinstance(raw, str):
                continue
            line = _clean_spaces(raw)
            if not line:
                continue

            for m in _ENTRY_RE.finditer(line):
                raw_nome = m.group("nome")
                faixa = _extrair_faixa(raw_nome)
                nome = _normalizar_nome(raw_nome)
                preco = _clean_spaces(m.group("preco")).replace("PÓ", "PO")
                if not nome or len(nome) < 3:
                    continue
                if "Magias" in nome or "Preço de Mercado" in nome:
                    continue
                if not _nome_valido(nome):
                    continue

                categoria, tipo = _inferir_categoria_tipo(ws.title, nome)
                if ws.title == "Tabela 7-17":
                    ln = nome.lower()
                    if "(poção" not in ln and "(óleo" not in ln:
                        continue
                tipos_alvo = ["Poção", "Óleo"] if tipo == "Poção/Óleo" else [tipo]
                for tipo_alvo in tipos_alvo:
                    chave = (nome.lower(), tipo_alvo.lower())
                    if chave in vistos:
                        continue
                    vistos.add(chave)

                    itens.append(
                        {
                            "nome": nome,
                            "descricao": f"Faixa d%: {faixa}" if faixa else None,
                            "pagina_referencia": (
                                "Livro do Mestre p.230"
                                if ws.title == "Tabela 7-17"
                                else "Livro do Mestre p.238-242"
                            ),
                            "categoria": categoria,
                            "tipo": tipo_alvo,
                            "custo": preco,
                            "peso": None,
                            "ativo": True,
                        }
                    )

    itens.sort(key=lambda x: (x["categoria"], x["tipo"], x["nome"]))
    return itens


def _py_repr_dict(d: dict) -> str:
    lines = ["    {"]
    for key, value in d.items():
        if value is None:
            lines.append(f"        {key!r}: None,")
        elif isinstance(value, bool):
            lines.append(f"        {key!r}: {value},")
        else:
            lines.append(f"        {key!r}: {value!r},")
    lines.append("    },")
    return "\n".join(lines)


def gerar_seed_script(itens: list[dict]) -> str:
    blocos = "\n".join(_py_repr_dict(i) for i in itens)
    return (
        '"""\n'
        "Seed de consumíveis D&D 3.5 (poções/óleos/pergaminhos).\n"
        "Gerado automaticamente por processar_consumiveis_excel.py.\n"
        '"""\n\n'
        "from datetime import datetime, timezone\n\n"
        "CONSUMIVEIS_DADOS = [\n"
        f"{blocos}\n"
        "]\n\n"
        r'''
def seed_consumiveis(db) -> None:
    from app.games.dnd35.models.consumivel import Consumivel
    from app.games.dnd35.models.consumivel import ConsumivelJogador

    nomes_seed = {d["nome"] for d in CONSUMIVEIS_DADOS}
    usados = {
        row[0]
        for row in db.query(ConsumivelJogador.consumivel_id).distinct().all()
        if row[0] is not None
    }
    legado = (
        db.query(Consumivel)
        .filter(Consumivel.deleted_at.is_(None))
        .all()
    )
    removidos = 0
    for item in legado:
        if item.nome not in nomes_seed and item.id not in usados:
            db.delete(item)
            removidos += 1
    if removidos:
        db.flush()

    inseridos = 0
    atualizados = 0
    for data in CONSUMIVEIS_DADOS:
        row = (
            db.query(Consumivel)
            .filter(Consumivel.nome == data["nome"], Consumivel.deleted_at.is_(None))
            .first()
        )
        campos = {
            "descricao": data.get("descricao"),
            "pagina_referencia": data.get("pagina_referencia"),
            "categoria": data.get("categoria"),
            "tipo": data.get("tipo"),
            "custo": data.get("custo"),
            "peso": data.get("peso"),
            "ativo": data.get("ativo", True),
        }
        if row:
            for k, v in campos.items():
                setattr(row, k, v)
            atualizados += 1
        else:
            db.add(
                Consumivel(
                    nome=data["nome"],
                    **campos,
                    criado_em=datetime.now(timezone.utc),
                )
            )
            inseridos += 1

    db.commit()
    print(
        f"✅ Catálogo de consumíveis sincronizado: +{inseridos} inseridos, "
        f"{atualizados} atualizados, {removidos} removidos (total: {len(CONSUMIVEIS_DADOS)})."
    )
'''
    )


if __name__ == "__main__":
    itens = extrair_consumiveis()
    if not itens:
        raise SystemExit("Nenhum consumível válido encontrado.")
    OUTPUT_PATH.parent.mkdir(parents=True, exist_ok=True)
    OUTPUT_PATH.write_text(gerar_seed_script(itens), encoding="utf-8")
    print(f"✅ Gerado: {OUTPUT_PATH} ({len(itens)} itens)")
