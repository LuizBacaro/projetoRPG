"""Service para importação em lote de magias via Excel (.xlsx/.xls)."""

from __future__ import annotations

from dataclasses import dataclass
from datetime import datetime, timedelta, timezone
from io import BytesIO
from threading import Lock
from typing import Any
from uuid import uuid4
import unicodedata

from fastapi import HTTPException, UploadFile

from ..schemas.magia import MagiaImportErro
from .magia_service import MagiaService


MAX_IMPORT_ROWS = 500
IMPORT_TTL_MINUTES = 30

CLASSES_VALIDAS = {
	"CLERIGO",
	"DRUIDA",
	"BARDO",
	"RANGER",
	"PALADINO",
	"MAGO",
	"FEITICEIRO",
}

ESCOLAS_VALIDAS = {
	"ABJURACAO",
	"CONJURACAO",
	"ADIVINHACAO",
	"ENCANTAMENTO",
	"EVOCACAO",
	"ILUSAO",
	"NECROMANCIA",
	"TRANSMUTACAO",
}

COMPONENTES_VALIDOS = {"V", "G", "M", "F", "FD", "XP", "S", "DF"}

HEADERS_ESPERADOS = [
	"nome",
	"nome_en",
	"escola",
	"sub_escola",
	"descritor",
	"classes_niveis",
	"componentes",
	"componente_extra",
	"tempo_conjuracao",
	"alcance",
	"area_efeito",
	"duracao",
	"teste_resistencia",
	"resistencia_magica",
	"resistencia_magia_texto",
	"dano",
	"descricao",
	"descricao_en",
	"e_magia_dominio",
	"dominios",
	"pagina_referencia",
]


@dataclass
class ImportDraft:
	rows: list[dict[str, Any]]
	expires_at: datetime


class MagiaImportService:
	_drafts: dict[str, ImportDraft] = {}
	_lock = Lock()

	def __init__(self, magia_service: MagiaService):
		self.magia_service = magia_service

	def gerar_modelo(self) -> bytes:
		try:
			from openpyxl import Workbook
		except Exception as exc:  # pragma: no cover
			raise HTTPException(status_code=500, detail="Dependência openpyxl não instalada") from exc

		wb = Workbook()
		ws = wb.active
		ws.title = "Magias"
		ws.append(HEADERS_ESPERADOS)
		ws.append(
			[
				"Misseis Magicos",
				"Magic Missile",
				"Evocacao",
				"",
				"Forca",
				"MAGO:1;FEITICEIRO:1",
				"V,G",
				"",
				"1 acao",
				"medio",
				"ate 5 projeteis",
				"instantanea",
				"Nenhum",
				"nao",
				"",
				"1d4+1 por projetil",
				"Cria dardos de energia magica.",
				"Creates darts of magical energy.",
				"nao",
				"",
				251,
			]
		)

		stream = BytesIO()
		wb.save(stream)
		stream.seek(0)
		return stream.read()

	async def criar_preview(self, file: UploadFile) -> dict[str, Any]:
		self._validar_arquivo_upload(file)
		content = await file.read()
		rows_raw = self._carregar_linhas_excel(file.filename or "", content)

		if len(rows_raw) > MAX_IMPORT_ROWS:
			raise HTTPException(
				status_code=422,
				detail=f"O arquivo excede o limite de {MAX_IMPORT_ROWS} magias por importação.",
			)

		erros: list[MagiaImportErro] = []
		validas: list[dict[str, Any]] = []
		for line_number, raw in rows_raw:
			item, erros_linha = self._normalizar_e_validar_linha(line_number, raw)
			if erros_linha:
				erros.extend(erros_linha)
				continue
			validas.append(item)

		import_id = self._armazenar_draft(validas)
		return {
			"import_id": import_id,
			"total_linhas": len(rows_raw),
			"validas": len(validas),
			"invalidas": len(erros),
			"preview": validas,
			"erros": [erro.model_dump() for erro in erros],
		}

	def confirmar_importacao(self, import_id: str) -> dict[str, Any]:
		rows = self._pop_draft(import_id)
		if rows is None:
			raise HTTPException(status_code=404, detail="Importação expirada ou inexistente")

		importadas = 0
		erros: list[MagiaImportErro] = []
		for idx, row in enumerate(rows, start=2):
			try:
				self.magia_service.criar(row)
				importadas += 1
			except HTTPException as exc:
				erros.append(
					MagiaImportErro(
						linha=idx,
						campo="geral",
						mensagem=str(exc.detail),
					)
				)

		return {
			"importadas": importadas,
			"falhas": len(erros),
			"erros": [erro.model_dump() for erro in erros],
		}

	def _validar_arquivo_upload(self, file: UploadFile) -> None:
		nome = (file.filename or "").lower()
		if not nome.endswith((".xlsx", ".xls")):
			raise HTTPException(status_code=422, detail="Formato de arquivo não suportado. Use .xlsx ou .xls")

	def _carregar_linhas_excel(self, filename: str, content: bytes) -> list[tuple[int, dict[str, Any]]]:
		lower_name = filename.lower()
		if lower_name.endswith(".xlsx"):
			return self._ler_xlsx(content)
		return self._ler_xls(content)

	def _ler_xlsx(self, content: bytes) -> list[tuple[int, dict[str, Any]]]:
		try:
			from openpyxl import load_workbook
		except Exception as exc:  # pragma: no cover
			raise HTTPException(status_code=500, detail="Dependência openpyxl não instalada") from exc

		wb = load_workbook(filename=BytesIO(content), data_only=True, read_only=True)
		ws = wb.active
		rows = list(ws.iter_rows(values_only=True))
		return self._rows_to_dicts(rows)

	def _ler_xls(self, content: bytes) -> list[tuple[int, dict[str, Any]]]:
		try:
			import xlrd
		except Exception as exc:  # pragma: no cover
			raise HTTPException(status_code=500, detail="Dependência xlrd não instalada") from exc

		wb = xlrd.open_workbook(file_contents=content)
		sheet = wb.sheet_by_index(0)
		rows = []
		for row_idx in range(sheet.nrows):
			rows.append(tuple(sheet.cell_value(row_idx, col_idx) for col_idx in range(sheet.ncols)))
		return self._rows_to_dicts(rows)

	def _rows_to_dicts(self, rows: list[tuple[Any, ...]]) -> list[tuple[int, dict[str, Any]]]:
		if not rows:
			raise HTTPException(status_code=422, detail="Arquivo vazio")

		headers = [self._slug(h) for h in rows[0]]
		if headers != HEADERS_ESPERADOS:
			raise HTTPException(
				status_code=422,
				detail="Cabeçalho inválido. Baixe o modelo de planilha e preencha as colunas esperadas.",
			)

		result: list[tuple[int, dict[str, Any]]] = []
		for i, row in enumerate(rows[1:], start=2):
			if self._linha_vazia(row):
				continue
			data = {headers[idx]: (row[idx] if idx < len(row) else None) for idx in range(len(headers))}
			result.append((i, data))
		return result

	@staticmethod
	def _linha_vazia(row: tuple[Any, ...]) -> bool:
		return all(str(item).strip() == "" for item in row if item is not None)

	def _normalizar_e_validar_linha(self, line_number: int, raw: dict[str, Any]) -> tuple[dict[str, Any] | None, list[MagiaImportErro]]:
		erros: list[MagiaImportErro] = []

		def val(key: str) -> str:
			value = raw.get(key)
			if value is None:
				return ""
			return str(value).strip()

		nome = val("nome")
		escola = val("escola")
		descricao = val("descricao")
		classes_niveis_raw = val("classes_niveis")

		if not nome:
			erros.append(self._erro(line_number, "nome", "Campo obrigatório"))
		if not escola:
			erros.append(self._erro(line_number, "escola", "Campo obrigatório"))
		if not descricao:
			erros.append(self._erro(line_number, "descricao", "Campo obrigatório"))

		escola_norm = self._normalizar_texto(escola)
		if escola and escola_norm not in ESCOLAS_VALIDAS:
			erros.append(self._erro(line_number, "escola", "Escola inválida"))

		classes_niveis = self._parse_classes_niveis(classes_niveis_raw, line_number, erros)
		componentes = self._normalizar_componentes(val("componentes"), line_number, erros)

		e_magia_dominio = self._to_bool(val("e_magia_dominio"))
		dominios = val("dominios") or None
		if e_magia_dominio:
			try:
				dominios = MagiaService.normalizar_dominios(dominios or "")
			except HTTPException as exc:
				erros.append(self._erro(line_number, "dominios", str(exc.detail)))
				dominios = None
		else:
			dominios = None

		pagina_referencia = self._to_int(val("pagina_referencia"), line_number, "pagina_referencia", erros)
		resistencia_magica = self._to_bool(val("resistencia_magica"))

		if erros:
			return None, erros

		item = {
			"nome": nome,
			"nome_en": val("nome_en") or None,
			"escola": escola,
			"sub_escola": val("sub_escola") or None,
			"descritor": val("descritor") or None,
			"componentes": componentes,
			"componente_extra": val("componente_extra") or None,
			"alcance": val("alcance") or None,
			"area_efeito": val("area_efeito") or None,
			"duracao": val("duracao") or None,
			"tempo_conjuracao": val("tempo_conjuracao") or None,
			"dano": val("dano") or None,
			"teste_resistencia": val("teste_resistencia") or None,
			"resistencia_magica": resistencia_magica,
			"resistencia_magia_texto": val("resistencia_magia_texto") or None,
			"descricao": descricao,
			"descricao_en": val("descricao_en") or None,
			"e_magia_dominio": e_magia_dominio,
			"dominios": dominios,
			"pagina_referencia": pagina_referencia,
			"classes_niveis": classes_niveis,
		}
		return item, []

	def _parse_classes_niveis(self, value: str, line_number: int, erros: list[MagiaImportErro]) -> list[dict[str, Any]]:
		if not value:
			erros.append(self._erro(line_number, "classes_niveis", "Informe ao menos uma classe com nível"))
			return []

		items = []
		vistos = set()
		for part in value.split(";"):
			trecho = part.strip()
			if not trecho:
				continue
			if ":" not in trecho:
				erros.append(self._erro(line_number, "classes_niveis", f"Formato inválido: {trecho}"))
				continue
			classe_raw, nivel_raw = trecho.split(":", 1)
			classe = self._normalizar_texto(classe_raw)
			if classe not in CLASSES_VALIDAS:
				erros.append(self._erro(line_number, "classes_niveis", f"Classe inválida: {classe_raw}"))
				continue
			if classe == "FEITICEIRO":
				classe = "MAGO"

			try:
				nivel = int(float(nivel_raw.strip()))
			except ValueError:
				erros.append(self._erro(line_number, "classes_niveis", f"Nível inválido para {classe_raw}"))
				continue
			if nivel < 0 or nivel > 9:
				erros.append(self._erro(line_number, "classes_niveis", f"Nível fora do intervalo 0-9 para {classe_raw}"))
				continue
			if classe in vistos:
				erros.append(self._erro(line_number, "classes_niveis", f"Classe duplicada: {classe}"))
				continue

			vistos.add(classe)
			items.append({"classe": classe, "nivel": nivel})

		if not items and not any(err.campo == "classes_niveis" for err in erros):
			erros.append(self._erro(line_number, "classes_niveis", "Informe ao menos uma classe com nível"))
		return items

	def _normalizar_componentes(self, value: str, line_number: int, erros: list[MagiaImportErro]) -> str | None:
		if not value:
			return None
		normalized = value.replace(" ", "").upper()
		normalized = normalized.replace("DF", "FD")
		normalized = normalized.replace("S", "G")
		parts = [part for part in normalized.split(",") if part]
		invalid = [part for part in parts if part not in COMPONENTES_VALIDOS]
		if invalid:
			erros.append(self._erro(line_number, "componentes", f"Componente inválido: {', '.join(invalid)}"))
			return None
		unique = []
		for comp in parts:
			if comp not in unique:
				unique.append(comp)
		return ",".join(unique)

	@staticmethod
	def _to_bool(value: str) -> bool:
		return value.strip().lower() in {"1", "true", "sim", "s", "yes", "y"}

	def _to_int(self, value: str, line_number: int, field: str, erros: list[MagiaImportErro]) -> int | None:
		if not value:
			return None
		try:
			parsed = int(float(value))
		except ValueError:
			erros.append(self._erro(line_number, field, "Número inválido"))
			return None
		if parsed <= 0:
			erros.append(self._erro(line_number, field, "Valor deve ser maior que zero"))
			return None
		return parsed

	@staticmethod
	def _erro(line_number: int, campo: str, mensagem: str) -> MagiaImportErro:
		return MagiaImportErro(linha=line_number, campo=campo, mensagem=mensagem)

	@staticmethod
	def _slug(value: Any) -> str:
		if value is None:
			return ""
		text = str(value).strip().lower()
		text = text.replace(" ", "_")
		return text

	@staticmethod
	def _normalizar_texto(value: str) -> str:
		normalized = unicodedata.normalize("NFD", value)
		sem_acentos = "".join(ch for ch in normalized if unicodedata.category(ch) != "Mn")
		return sem_acentos.strip().upper()

	def _armazenar_draft(self, rows: list[dict[str, Any]]) -> str:
		self._limpar_expirados()
		import_id = uuid4().hex
		with self._lock:
			self._drafts[import_id] = ImportDraft(
				rows=rows,
				expires_at=datetime.now(timezone.utc) + timedelta(minutes=IMPORT_TTL_MINUTES),
			)
		return import_id

	def _pop_draft(self, import_id: str) -> list[dict[str, Any]] | None:
		self._limpar_expirados()
		with self._lock:
			draft = self._drafts.pop(import_id, None)
		if not draft:
			return None
		return draft.rows

	def _limpar_expirados(self) -> None:
		now = datetime.now(timezone.utc)
		with self._lock:
			expirados = [key for key, draft in self._drafts.items() if draft.expires_at <= now]
			for key in expirados:
				self._drafts.pop(key, None)
