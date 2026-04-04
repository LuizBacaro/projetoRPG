"""
Loader de Perícias do Excel: Perícias.xlsx
Responsável por ler e estruturar dados de perícias do arquivo Excel
"""

import openpyxl
import unicodedata
from typing import List, Dict, Set
from pathlib import Path


def remover_acentos(texto: str) -> str:
    """Remove acentuação de um string"""
    if not texto:
        return texto
    nfd = unicodedata.normalize('NFD', texto)
    return ''.join(char for char in nfd if unicodedata.category(char) != 'Mn')


class PericiaExcelLoader:
    """
    Carrega e valida perícias do arquivo Excel
    Estrutura esperada: Colunas com perícia, descrição, classe, especialidade, etc.
    """
    
    CLASSES_D_D = {
        'Barbaro', 'Bardo', 'Clerico', 'Druida', 'Feiticeiro', 
        'Guerreiro', 'Ladino', 'Mago', 'Monge', 'Paladino', 'Ranger'
    }
    
    ATRIBUTOS = {'FOR', 'DES', 'CON', 'INT', 'SAB', 'CAR'}
    
    def __init__(self, caminho_excel: str):
        """Inicializa o loader com caminho do arquivo Excel"""
        self.caminho = Path(caminho_excel)
        if not self.caminho.exists():
            raise FileNotFoundError(f"Arquivo não encontrado: {caminho_excel}")
        
        self.wb = openpyxl.load_workbook(self.caminho)
        self.ws = self.wb['Perícias']
        self.pericias_data = []
        self.pericia_classe_map = {}
    
    def carregar(self) -> Dict:
        """
        Carrega todas as perícias do Excel
        Retorna: {
            'pericias': [lista de perícias],
            'pericia_classe': {pericia_nome: {classe: True/False}}
        }
        """
        # Pegar headers na linha 2
        headers = self._extrair_headers()
        
        # Processar linhas a partir da linha 5 (linha 4 tem mais headers)
        for row_idx in range(5, self.ws.max_row + 1):
            row_data = self._extrair_linha(row_idx, headers)
            if row_data and row_data.get('nome'):
                self.pericias_data.append(row_data)
        
        return {
            'pericias': self.pericias_data,
            'pericia_classe': self.pericia_classe_map
        }
    
    def _extrair_headers(self) -> Dict[str, int]:
        """
        Extrai mapeamento coluna -> índice dos headers
        Esperado na linha 2
        """
        headers = {}
        row2 = list(self.ws.iter_rows(min_row=2, max_row=2, values_only=True))[0]
        
        for idx, valor in enumerate(row2):
            if not valor:
                continue
            
            valor_lower = str(valor).lower().strip()
            
            # Mapear headers conhecidos
            if 'perícia' in valor_lower and 'especialização' not in valor_lower:
                headers['nome'] = idx
            elif 'descrição' in valor_lower:
                headers['descricao'] = idx
            elif 'requer especialização' in valor_lower:
                headers['requer_especialidade'] = idx
            elif 'especialização' in valor_lower and 'associação' not in valor_lower:
                headers['especialidade'] = idx
            elif 'atributo base' in valor_lower or 'atributo' in valor_lower:
                headers['atributo'] = idx
            elif 'penalidade de armadura' in valor_lower:
                headers['sofre_penalidade_armadura'] = idx
            elif 'sem treinamento' in valor_lower or 'pode ser utilizada' in valor_lower:
                headers['pode_usar_sem_treinamento'] = idx
            elif 'associação' in valor_lower and 'classe' in valor_lower:
                # Marca início da seção de classes
                headers['classes_start'] = idx
        
        # Extrair nomes das classes da linha 3
        row3 = list(self.ws.iter_rows(min_row=3, max_row=3, values_only=True))[0]
        headers['classes_map'] = {}  # {coluna_idx: nome_classe}
        
        for idx, valor in enumerate(row3):
            if valor:
                classe_normalizada = remover_acentos(str(valor).strip())
                if classe_normalizada in self.CLASSES_D_D:
                    headers['classes_map'][idx] = classe_normalizada
        
        return headers
    
    def _extrair_linha(self, row_idx: int, headers: Dict) -> Dict:
        """Extrai dados de uma linha do Excel"""
        row_values = list(self.ws.iter_rows(
            min_row=row_idx, 
            max_row=row_idx, 
            values_only=True
        ))[0]
        
        # Extrair campos básicos
        nome_idx = headers.get('nome', 2)
        nome = row_values[nome_idx] if nome_idx is not None and nome_idx < len(row_values) else None
        
        if not nome or str(nome).strip() == '':
            return None
        
        nome = str(nome).strip()
        
        pericia = {
            'id': int(row_values[1]) if row_values[1] else None,
            'nome': nome,
            'descricao': str(row_values[headers.get('descricao', 3)]).strip() if headers.get('descricao') is not None else '',
            'atributo': self._normalizar_atributo(row_values[headers.get('atributo', 22)]),
            'especialidade': self._limpar_valor(row_values[headers.get('especialidade', 5)]),
            'requer_treinamento': 1 if str(row_values[headers.get('requer_especialidade', 4)]).upper() == 'SIM' else 0,
            'pode_usar_sem_treinamento': 1 if str(row_values[headers.get('pode_usar_sem_treinamento', 21)]).upper() == 'SIM' else 0,
            'sofre_penalidade_armadura': 1 if str(row_values[headers.get('sofre_penalidade_armadura', 23)]).upper() == 'SIM' else 0,
            'tipo': 'comum',  # Default
        }
        
        # Determinar tipo baseado em padrões
        if 'Conhecimento:' in nome:
            pericia['tipo'] = 'conhecimento'
        elif 'Atuação' in nome or 'Profissão' in nome:
            pericia['tipo'] = 'performance' if 'Atuação' in nome else 'profissao'
        
        # Extrair associações com classes usando o mapa
        classes_set = set()
        if 'classes_map' in headers:
            for col_idx, classe_nome in headers['classes_map'].items():
                if col_idx < len(row_values):
                    valor = row_values[col_idx]
                    if valor and str(valor).strip() == 'X':
                        classes_set.add(classe_nome)
        
        # Guardar mapeamento classe
        if nome not in self.pericia_classe_map:
            self.pericia_classe_map[nome] = {}
        
        for classe in self.CLASSES_D_D:
            self.pericia_classe_map[nome][classe] = (classe in classes_set)
        
        return pericia
    
    def _normalizar_atributo(self, valor) -> str:
        """Normaliza valor de atributo para padrão"""
        if not valor:
            return 'INT'
        
        valor_str = str(valor).upper().strip()
        
        # Mapeamentos
        mapa = {
            'FOR': 'FOR',
            'FORÇA': 'FOR',
            'DES': 'DES',
            'DESTREZA': 'DES',
            'CON': 'CON',
            'CONSTITUIÇÃO': 'CON',
            'INT': 'INT',
            'INTELIGÊNCIA': 'INT',
            'SAB': 'SAB',
            'SABEDORIA': 'SAB',
            'CAR': 'CAR',
            'CARISMA': 'CAR',
        }
        
        return mapa.get(valor_str, 'INT')
    
    def _limpar_valor(self, valor) -> str:
        """Limpa e normaliza valor de string"""
        if not valor or str(valor).strip() == '':
            return None
        return str(valor).strip()
    
    def validar(self) -> List[str]:
        """Valida dados carregados, retorna lista de erros"""
        erros = []
        
        for pericia in self.pericias_data:
            # Validar atributo
            if pericia['atributo'] not in self.ATRIBUTOS:
                erros.append(f"{pericia['nome']}: Atributo inválido '{pericia['atributo']}'")
            
            # Validar tipo
            tipos_validos = {'comum', 'conhecimento', 'profissao', 'oficio', 'performance'}
            if pericia['tipo'] not in tipos_validos:
                erros.append(f"{pericia['nome']}: Tipo inválido '{pericia['tipo']}'")
            
            # Validar se tem pelo menos uma classe associada
            classes_associadas = self.pericia_classe_map.get(pericia['nome'], {})
            if not any(classes_associadas.values()):
                erros.append(f"{pericia['nome']}: Nenhuma classe associada")
        
        return erros
    
    def gerar_seed_data(self) -> str:
        """Gera código Python para seed baseado nos dados carregados"""
        linhas = [
            "# Gerado automaticamente por pericias_loader.py",
            "PERICIAS_DO_EXCEL = [",
        ]
        
        for p in self.pericias_data:
            descricao_escaped = p['descricao'].replace('"', '\\"')
            bloco = f"""    {{
        "nome": "{p['nome']}",
        "descricao": "{descricao_escaped}",
        "atributo": "{p['atributo']}",
        "tipo": "{p['tipo']}",
        "especialidade": {repr(p['especialidade'])},
        "requer_treinamento": {p['requer_treinamento']},
        "pode_usar_sem_treinamento": {p['pode_usar_sem_treinamento']},
        "sofre_penalidade_armadura": {p['sofre_penalidade_armadura']},
    }},"""
            linhas.append(bloco)
        
        linhas.append("]")
        linhas.append("")
        linhas.append("PERICIA_CLASSES = {")
        
        for pericia, classes_dict in self.pericia_classe_map.items():
            classes_default = [c for c, é_default in classes_dict.items() if é_default]
            linhas.append(f'    "{pericia}": {classes_default},')
        
        linhas.append("}")
        
        return "\n".join(linhas)


def carregar_pericias(caminho_excel: str) -> Dict:
    """
    Função simples para carregar péricias do Excel
    
    Uso:
        pericias_data = carregar_pericias('/caminho/para/Perícias.xlsx')
        pericias = pericias_data['pericias']
        mapa_classes = pericias_data['pericia_classe']
    """
    loader = PericiaExcelLoader(caminho_excel)
    dados = loader.carregar()
    
    # Validar
    erros = loader.validar()
    if erros:
        print("⚠️ Erros de validação encontrados:")
        for erro in erros:
            print(f"  - {erro}")
    
    return dados


if __name__ == '__main__':
    # Teste direto
    import sys
    
    caminho = '/home/luiz/gitHub/projetoRPG/Perícias.xlsx'
    
    print("📚 Carregando Perícias do Excel...")
    loader = PericiaExcelLoader(caminho)
    dados = loader.carregar()
    
    print(f"\n✅ {len(dados['pericias'])} perícias carregadas!")
    print(f"✅ {sum(sum(v.values()) for v in dados['pericia_classe'].values())} associações classe-perícia!")
    
    # Mostrar amostra
    print("\n📋 Primeiras 5 perícias:")
    for p in dados['pericias'][:5]:
        print(f"  - {p['nome']} ({p['atributo']}) - Tipo: {p['tipo']}")
    
    # Validar
    erros = loader.validar()
    if erros:
        print(f"\n⚠️ {len(erros)} erros de validação:")
        for erro in erros[:5]:
            print(f"  - {erro}")
    else:
        print("\n✅ Todas as perícias validadas com sucesso!")
    
    # Gerar preview do seed
    print("\n🔧 Preview do código de seed:")
    seed_code = loader.gerar_seed_data()
    print(seed_code[:500] + "...")
