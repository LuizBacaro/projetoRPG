"""
Seed simplificado - popula banco de dados
"""
import sys
from pathlib import Path
import openpyxl  # ✅ Adicionado

# Adicionar diretório raiz ao path
backend_dir = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(backend_dir))

print(f"📁 Backend dir: {backend_dir}")

try:
    from app.core.database import SessionLocal, Base, engine
    from app.models.combatente import Combatente
    import app.models.combate  # Importar para criar tabela
    
    print("✅ Imports bem-sucedidos!")
    
    # Criar tabelas
    print("🔧 Criando tabelas...")
    Base.metadata.create_all(bind=engine)
    print("✅ Tabelas criadas!")
    
    # Popular banco
    db = SessionLocal()
    
    try:
        # Verificar se já tem dados
        count = db.query(Combatente).count()
        print(f"📊 Combatentes existentes: {count}")
        
        if count > 0:
            print("\n⚠️  Banco já possui dados!")
            print("Continuando sem limpar (mantenha os dados existentes)")
            print("Se quiser resetar, delete o arquivo rpg_arena.db e rode novamente\n")
            
            # Verificar e popular perícias se vazio
            from app.models.pericia import PericiaClasse, Pericia
            
            pericia_count = db.query(Pericia).count()
            if pericia_count == 0:
                print("🔄 Populando perícias do Excel...")
                excel_path = backend_dir.parent / 'Perícias.xlsx'
                
                if excel_path.exists():
                    wb = openpyxl.load_workbook(excel_path)
                    ws = wb.active
                    
                    # Row 2 has headers
                    row2 = list(ws.iter_rows(min_row=2, max_row=2, values_only=True))[0]
                    headers = {}
                    for idx, valor in enumerate(row2):
                        if not valor:
                            continue
                        valor_lower = str(valor).lower().strip()
                        if 'perícia' in valor_lower and 'especialização' not in valor_lower:
                            headers['nome'] = idx
                        elif 'descrição' in valor_lower:
                            headers['descricao'] = idx
                        elif 'atributo' in valor_lower:
                            headers['atributo'] = idx
                        elif 'penalidade de armadura' in valor_lower:
                            headers['sofre_penalidade_armadura'] = idx
                        elif 'sem treinamento' in valor_lower or 'pode ser utilizada' in valor_lower:
                            headers['pode_usar_sem_treinamento'] = idx
                    
                    # Process data rows
                    pericias_list = []
                    for row_idx in range(5, ws.max_row + 1):
                        row = list(ws.iter_rows(min_row=row_idx, max_row=row_idx, values_only=True))[0]
                        
                        nome = row[headers.get('nome', 2)] if headers.get('nome') and len(row) > headers.get('nome') else None
                        if not nome or not isinstance(nome, str):
                            continue
                        
                        nome = nome.strip()
                        if not nome:
                            continue
                        
                        # Check if already exists
                        if db.query(Pericia).filter(Pericia.nome == nome).first():
                            continue
                        
                        # Get field values safely
                        desc = row[headers.get('descricao', 3)] if headers.get('descricao') and len(row) > headers.get('descricao') else ''
                        attr = row[headers.get('atributo', -1)] if headers.get('atributo') and len(row) > headers.get('atributo') else 'DES'
                        pode_usar = 1 if (headers.get('pode_usar_sem_treinamento') and len(row) > headers.get('pode_usar_sem_treinamento') and row[headers.get('pode_usar_sem_treinamento')] == 'Sim') else 0
                        sofre_pen = 1 if (headers.get('sofre_penalidade_armadura') and len(row) > headers.get('sofre_penalidade_armadura') and row[headers.get('sofre_penalidade_armadura')] == 'Sim') else 0
                        
                        pericia = Pericia(
                            nome=nome,
                            descricao=desc,
                            atributo=attr,
                            tipo='comum',
                            requer_treinamento=0,
                            especialidade=None,
                            pode_usar_sem_treinamento=pode_usar,
                            sofre_penalidade_armadura=sofre_pen,
                        )
                        pericias_list.append(pericia)
                    
                    # Add all
                    db.add_all(pericias_list)
                    db.commit()
                    print(f"✅ {len(pericias_list)} perícias criadas!")
                else:
                    print(f"⚠️  Excel não encontrado: {excel_path}")
            else:
                print(f"✅ Perícias já populadas ({pericia_count} perícias)")
            
            # Verificar e popular pericias_classes se vazio
            pericia_class_count = db.query(PericiaClasse).count()
            if pericia_class_count == 0:
                print("🔄 Populando pericias_classes do Excel...")
                excel_path = backend_dir.parent / 'Perícias.xlsx'
                
                if excel_path.exists():
                    wb = openpyxl.load_workbook(excel_path)
                    ws = wb.active
                    
                    # Row 3 tem class names
                    row3 = list(ws.iter_rows(min_row=3, max_row=3, values_only=True))[0]
                    classe_columns = {}
                    
                    for col_idx in range(6, len(row3)):
                        classe = row3[col_idx]
                        if classe and isinstance(classe, str) and classe.strip() and classe.strip() not in ['Todos', '']:
                            classe_columns[col_idx] = classe.strip()
                    
                    # Process data rows
                    associacoes = []
                    for row_idx in range(5, ws.max_row + 1):
                        row = list(ws.iter_rows(min_row=row_idx, max_row=row_idx, values_only=True))[0]
                        
                        nome_pericia = row[2] if len(row) > 2 else None
                        if not nome_pericia or not isinstance(nome_pericia, str):
                            continue
                        
                        nome_pericia = nome_pericia.strip()
                        if not nome_pericia:
                            continue
                        
                        pericia = db.query(Pericia).filter(Pericia.nome == nome_pericia).first()
                        if not pericia:
                            continue
                        
                        # Check columns
                        for col_idx, classe_nome in classe_columns.items():
                            cell_value = row[col_idx] if len(row) > col_idx else None
                            if cell_value == 'X':
                                assoc = PericiaClasse(
                                    pericia_id=pericia.id,
                                    classe_nome=classe_nome,
                                    is_default=1
                                )
                                associacoes.append(assoc)
                    
                    # Add all
                    db.add_all(associacoes)
                    db.commit()
                    print(f"✅ {len(associacoes)} associações pericias_classes criadas!")
                else:
                    print(f"⚠️  Excel não encontrado: {excel_path}")
            else:
                print(f"✅ pericias_classes já populado ({pericia_class_count} associações)")
        
        else:
            print("\n🌱 Populando banco de dados...")
            
            combatentes_iniciais = [
                {
                    "nome": "Theron", "tipo": "jogador", "classe": "Guerreiro",
                    "hp_maximo": 85, "hp_atual": 85, "iniciativa": 15,
                    "forca": 16, "destreza": 12, "constituicao": 14,
                    "inteligencia": 10, "sabedoria": 11, "carisma": 13,
                    "nivel": 5, "pontos": 1200
                },
                {
                    "nome": "Lyra", "tipo": "jogador", "classe": "Mago",
                    "hp_maximo": 45, "hp_atual": 45, "iniciativa": 18,
                    "forca": 8, "destreza": 14, "constituicao": 10,
                    "inteligencia": 18, "sabedoria": 15, "carisma": 12,
                    "nivel": 5, "pontos": 1150
                },
                {
                    "nome": "Garrick", "tipo": "jogador", "classe": "Clérigo",
                    "hp_maximo": 65, "hp_atual": 65, "iniciativa": 12,
                    "forca": 14, "destreza": 10, "constituicao": 13,
                    "inteligencia": 12, "sabedoria": 16, "carisma": 14,
                    "nivel": 5, "pontos": 980
                },
                {
                    "nome": "Zara", "tipo": "jogador", "classe": "Ladino",
                    "hp_maximo": 55, "hp_atual": 55, "iniciativa": 20,
                    "forca": 10, "destreza": 18, "constituicao": 12,
                    "inteligencia": 14, "sabedoria": 13, "carisma": 15,
                    "nivel": 5, "pontos": 1350
                },
                {
                    "nome": "Goblin Arqueiro", "tipo": "monstro", "classe": "Arqueiro",
                    "hp_maximo": 30, "hp_atual": 30, "iniciativa": 14,
                    "forca": 8, "destreza": 14, "constituicao": 10,
                    "inteligencia": 10, "sabedoria": 8, "carisma": 8,
                    "nivel": 2, "pontos": 0
                },
                {
                    "nome": "Orc Guerreiro", "tipo": "monstro", "classe": "Guerreiro",
                    "hp_maximo": 60, "hp_atual": 60, "iniciativa": 10,
                    "forca": 16, "destreza": 12, "constituicao": 16,
                    "inteligencia": 7, "sabedoria": 11, "carisma": 10,
                    "nivel": 3, "pontos": 0
                },
                {
                    "nome": "Comerciante Anão", "tipo": "npc", "classe": "Civil",
                    "hp_maximo": 40, "hp_atual": 40, "iniciativa": 8,
                    "forca": 12, "destreza": 10, "constituicao": 14,
                    "inteligencia": 12, "sabedoria": 13, "carisma": 11,
                    "nivel": 1, "pontos": 0
                }
            ]
            
            for data in combatentes_iniciais:
                combatente = Combatente(**data)
                db.add(combatente)
                print(f"  ✅ {data['nome']} ({data['tipo']}) - Nível {data['nivel']}")
            
            db.commit()
            print(f"\n🎉 {len(combatentes_iniciais)} combatentes criados com sucesso!")
        
        print("\n" + "=" * 50)
        print("Para iniciar o servidor:")
        print("  python -m uvicorn app.main:app --reload")
        print("=" * 50)
        
    finally:
        db.close()
        
except ImportError as e:
    print(f"\n❌ Erro de importação: {e}")
    print("\n🔍 Verificações necessárias:")
    print(f"  1. Verifique se existe: {backend_dir}/app/models/combatente.py")
    print(f"  2. Verifique se existe: {backend_dir}/app/models/__init__.py")
    print(f"  3. Verifique se existe: {backend_dir}/app/__init__.py")
    print(f"  4. Verifique se existe: {backend_dir}/app/core/database.py")
    print("\nExecute para diagnosticar:")
    print("  find app -name '*.py' | head -20")
    sys.exit(1)
    
except Exception as e:
    print(f"\n❌ Erro: {e}")
    import traceback
    traceback.print_exc()
    sys.exit(1)