"""
Gera catálogo racial normalizado a partir da planilha
"Características especiais.xlsx".
"""

from __future__ import annotations

import json
import re
import unicodedata
from datetime import datetime, timezone
from pathlib import Path
from typing import Any

from openpyxl import load_workbook


PLANILHA_PATH = Path("Características especiais.xlsx")
OUTPUT_PATH = Path("docs/dados/racas_caracteristicas_catalogo.json")

_ATTR_ALIASES = {
    "FORCA": "forca",
    "DESTREZA": "destreza",
    "CONSTITUICAO": "constituicao",
    "INTELIGENCIA": "inteligencia",
    "SABEDORIA": "sabedoria",
    "CARISMA": "carisma",
}


def _sem_acentos(valor: str) -> str:
    normal = unicodedata.normalize("NFD", valor)
    return "".join(ch for ch in normal if unicodedata.category(ch) != "Mn")


def _slug(valor: str) -> str:
    base = _sem_acentos(str(valor or "")).lower().strip()
    base = re.sub(r"[^a-z0-9]+", "-", base)
    return base.strip("-")


def _texto_limpo(valor: Any) -> str:
    if valor is None:
        return ""
    texto = str(valor).strip()
    if texto.lower() in {"nan", "none"}:
        return ""
    return texto


def _lista_linhas(valor: Any) -> list[str]:
    texto = _texto_limpo(valor)
    if not texto or texto == "-":
        return []
    linhas = []
    for parte in texto.replace("\r", "\n").split("\n"):
        item = parte.strip()
        if not item:
            continue
        if item.startswith("-"):
            item = item[1:].strip()
        if item:
            linhas.append(item)
    return linhas


def _lista_separada_por_pontoevirgula(valor: Any) -> list[str]:
    texto = _texto_limpo(valor)
    if not texto or texto == "-":
        return []
    return [p.strip() for p in texto.split(";") if p.strip()]


def _parse_deslocamento_metros(valor: Any) -> int | None:
    texto = _texto_limpo(valor).lower()
    if not texto:
        return None
    m = re.search(r"(\d+)", texto)
    return int(m.group(1)) if m else None


def _parse_modificadores_habilidade(valor: Any) -> list[dict[str, Any]]:
    texto = _texto_limpo(valor)
    if not texto or texto == "-":
        return []
    normal = _sem_acentos(texto).upper()
    tokens = re.split(r"[;,]", normal)
    saida: list[dict[str, Any]] = []
    for token in tokens:
        item = token.strip()
        if not item:
            continue
        m = re.match(r"([+-]\d+)\s+([A-Z ]+)$", item)
        if not m:
            continue
        valor_num = int(m.group(1))
        attr_raw = re.sub(r"\s+", " ", m.group(2).strip())
        attr = _ATTR_ALIASES.get(attr_raw)
        if not attr:
            continue
        saida.append(
            {
                "atributo": attr,
                "valor": valor_num,
            }
        )
    return saida


def processar_planilha() -> dict[str, Any]:
    if not PLANILHA_PATH.exists():
        raise FileNotFoundError(f"Arquivo não encontrado: {PLANILHA_PATH}")

    wb = load_workbook(PLANILHA_PATH, data_only=True)
    if "Raças" not in wb.sheetnames:
        raise ValueError("Aba 'Raças' não encontrada na planilha.")

    ws = wb["Raças"]
    headers_row = 3
    headers = [_texto_limpo(c.value) for c in ws[headers_row]]
    idx = {nome: i for i, nome in enumerate(headers) if nome}

    racas: list[dict[str, Any]] = []
    for row_number in range(headers_row + 1, ws.max_row + 1):
        row = [c.value for c in ws[row_number]]
        nome = _texto_limpo(row[idx["Raça"]]) if "Raça" in idx else ""
        if not nome:
            continue

        registro = {
            "slug": _slug(nome),
            "nome": nome,
            "modificadores_habilidade": _parse_modificadores_habilidade(
                row[idx.get("Modificadores de habilidades", -1)] if "Modificadores de habilidades" in idx else ""
            ),
            "tamanho": _texto_limpo(row[idx.get("Tamanho", -1)]) if "Tamanho" in idx else "",
            "deslocamento_metros": _parse_deslocamento_metros(
                row[idx.get("Deslocamento", -1)] if "Deslocamento" in idx else ""
            ),
            "idiomas_iniciais": _lista_separada_por_pontoevirgula(
                row[idx.get("Idiomas iniciais", -1)] if "Idiomas iniciais" in idx else ""
            ),
            "talentos_especiais": _lista_linhas(
                row[idx.get("Talentos especiais", -1)] if "Talentos especiais" in idx else ""
            ),
            "habilidades_especiais": _lista_linhas(
                row[idx.get("Habilidades especiais", -1)] if "Habilidades especiais" in idx else ""
            ),
            "resistencias": _lista_linhas(
                row[idx.get("Resistências", -1)] if "Resistências" in idx else ""
            ),
            "modificadores_ataque": _lista_linhas(
                row[idx.get("Modificadores de ataque", -1)] if "Modificadores de ataque" in idx else ""
            ),
            "modificadores_defesa": _lista_linhas(
                row[idx.get("Modificadores de defesa", -1)] if "Modificadores de defesa" in idx else ""
            ),
            "modificadores_pericia": _lista_linhas(
                row[idx.get("Modificadores de perícia", -1)] if "Modificadores de perícia" in idx else ""
            ),
            "classe_favorecida": _texto_limpo(
                row[idx.get("Classe favorecida", -1)] if "Classe favorecida" in idx else ""
            ),
        }
        racas.append(registro)

    return {
        "source": str(PLANILHA_PATH),
        "generated_at": datetime.now(timezone.utc).isoformat(),
        "total_racas": len(racas),
        "racas": racas,
    }


def main() -> int:
    payload = processar_planilha()
    OUTPUT_PATH.parent.mkdir(parents=True, exist_ok=True)
    OUTPUT_PATH.write_text(json.dumps(payload, ensure_ascii=False, indent=2), encoding="utf-8")
    print(f"✅ Catálogo racial gerado: {OUTPUT_PATH} ({payload['total_racas']} raças)")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
